"""标准教师安排模板导入的公开 API 契约测试。"""

import io
import json
from datetime import time

import pytest
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError

from app.api.imports import XLSX_MIME
from app.models.assignment import CourseAssignment, SchedulingUnit
from app.models.audit import AuditLog
from app.models.basedata import ClassUnit, Room, Subject, Teacher
from app.models.period import Period, PeriodTable, PeriodType
from app.models.semester import Semester
from app.models.teacher_arrangement_import import (
    TeacherArrangementImportBatch,
    TeacherArrangementImportRecord,
    TeacherArrangementSourceRecord,
)
from app.models.timetable import Timetable
from app.models.user import Role
from app.services import teacher_arrangement_import
from tests.conftest import make_user

PW = "password123"


@pytest.fixture
def import_env(env):
    client, db = env
    make_user(db, "arrangement-director", PW, roles=[Role.director])
    client.post(
        "/api/auth/login",
        json={"username": "arrangement-director", "password": PW},
    )
    response = client.post(
        "/api/semesters",
        json={
            "academic_year": 2026,
            "term": 1,
            "start_date": "2026-09-01",
            "end_date": "2027-01-20",
        },
    )
    assert response.status_code == 201, response.json()
    return client, db, response.json()["id"]


def set_row(workbook, sheet_name: str, values: dict[str, object], row: int = 4) -> None:
    sheet = workbook[sheet_name]
    columns = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
    for header, value in values.items():
        sheet.cell(row=row, column=columns[header]).value = value


def clear_data_row(workbook, sheet_name: str, row: int = 4) -> None:
    sheet = workbook[sheet_name]
    for column in range(1, sheet.max_column + 1):
        sheet.cell(row=row, column=column).value = None


def minimal_standard_workbook(
    client,
    semester_id: int,
    *,
    suffix: str = "",
    mode: str = "standard",
) -> bytes:
    response = client.get(
        "/api/import/teacher-arrangements/template",
        params={"semester_id": semester_id, "mode": mode},
    )
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    set_row(
        workbook,
        "科目",
        {
            "学校科目编码": "SUB-MATH",
            "科目名称": f"数学{suffix}",
            "领域/类别": "数学",
            "所需教室/场地类型": "普通教室",
        },
    )
    set_row(
        workbook,
        "教师",
        {
            "学校教师编码": "T-001",
            "教师姓名": f"王老师{suffix}",
            "基础周课时": 18,
            "行政职务": "年级负责人",
            "行政减课时": 2,
            "教师状态": "在岗",
            "外聘": "否",
        },
    )
    class_values: dict[str, object] = {
        "学校班级编码": "C-701",
        "班级名称": f"七年级1班{suffix}",
        "年级": 7,
        "学制": "初中",
        "专业/班级类别": "普通班",
        "班主任": "T-001",
        "班级计划周课时": 5,
    }
    if mode == "scheduling_ready":
        class_values["作息表编码"] = "PT-JUNIOR"
    set_row(
        workbook,
        "班级",
        class_values,
    )
    set_row(
        workbook,
        "教学任务",
        {
            "任务编码": "TASK-701-MATH",
            "班级": "C-701",
            "科目": "SUB-MATH",
            "组成": "基础课",
            "周课时": 5,
            "主讲教师": "T-001",
            "协同教师": "",
        },
    )
    set_row(
        workbook,
        "来源记录",
        {
            "记录编码": "SRC-001",
            "记录类型": "备注",
            "关联任务编码": "TASK-701-MATH",
            "原始内容": "人数与费用由年级另行维护",
            "备注": "来自教师安排表备注栏",
        },
    )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def complete_ready_workbook(client, semester_id: int) -> bytes:
    content = minimal_standard_workbook(
        client,
        semester_id,
        mode="scheduling_ready",
    )
    workbook = load_workbook(io.BytesIO(content))
    set_row(
        workbook,
        "科目",
        {"所需教室/场地类型": "专用教室"},
    )
    set_row(
        workbook,
        "教室及户外场地",
        {
            "学校教室/场地编码": "ROOM-MATH",
            "教室/场地名称": "数学专用教室",
            "教室/场地类型": "专用教室",
            "容量": 48,
            "适用科目": "SUB-MATH",
        },
    )
    for offset, weekday in enumerate(
        ("星期一", "星期二", "星期三", "星期四", "星期五")
    ):
        set_row(
            workbook,
            "作息时间表",
            {
                "作息表编码": "PT-JUNIOR",
                "作息表名称": "初中部作息",
                "星期": weekday,
                "节次": 1,
                "课时类型": "常规课时",
                "开始时间": time(8, 0),
                "结束时间": time(8, 45),
            },
            row=4 + offset,
        )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def post_workbook(
    client,
    action: str,
    semester_id: int,
    content: bytes,
    data: dict[str, str] | None = None,
    *,
    mode: str = "standard",
):
    return client.post(
        f"/api/import/teacher-arrangements/{action}",
        params={"semester_id": semester_id, "mode": mode},
        data=data,
        files={"file": ("teacher-arrangement.xlsx", content, XLSX_MIME)},
    )


def test_standard_preview_is_zero_write_and_commit_creates_scheduling_data(import_env):
    client, db, semester_id = import_env
    content = minimal_standard_workbook(client, semester_id)

    response = post_workbook(client, "preview", semester_id, content)

    assert response.status_code == 200, response.json()
    preview = response.json()
    assert preview["template_version"] == "1.0"
    assert preview["mode"] == "standard"
    assert preview["can_commit"] is True
    assert preview["counts"] == {
        "new": 5,
        "changed": 0,
        "unchanged": 0,
        "conflict": 0,
        "disappeared": 0,
        "blocker": 0,
        "warning": 0,
    }
    assert preview["issues"] == []
    assert db.query(Subject).count() == 0
    assert db.query(Teacher).count() == 0
    assert db.query(ClassUnit).count() == 0
    assert db.query(CourseAssignment).count() == 0

    committed = post_workbook(
        client,
        "commit",
        semester_id,
        content,
        data={"fingerprint": preview["fingerprint"], "confirm_changes": "false"},
    )

    assert committed.status_code == 200, committed.json()
    result = committed.json()
    assert result["idempotent"] is False
    assert result["created"] == {
        "subjects": 1,
        "teachers": 1,
        "classes": 1,
        "assignments": 1,
        "source_records": 1,
    }
    subject = db.query(Subject).one()
    teacher = db.query(Teacher).one()
    class_unit = db.query(ClassUnit).one()
    assignment = db.query(CourseAssignment).one()
    source = db.query(TeacherArrangementSourceRecord).one()
    semester = db.get(Semester, semester_id)
    assert subject.school_code == "SUB-MATH"
    assert teacher.school_code == "T-001"
    assert teacher.arrangement_status == "normal"
    assert class_unit.school_code == "C-701"
    assert class_unit.planned_weekly_periods == 5
    assert class_unit.homeroom_teacher_id == teacher.id
    assert assignment.task_code == "TASK-701-MATH"
    assert assignment.component == "基础课"
    assert assignment.subject_id == subject.id
    assert assignment.scheduling_unit.members[0].class_unit_id == class_unit.id
    assert [(link.teacher_id, link.is_lead) for link in assignment.teachers] == [
        (teacher.id, True)
    ]
    assert source.content == "人数与费用由年级另行维护"
    assert semester is not None and semester.readiness == "draft"
    assert db.query(SchedulingUnit).count() == 1
    assert db.query(Timetable).count() == 0
    assert db.query(TeacherArrangementImportBatch).count() == 1
    assert db.query(TeacherArrangementImportRecord).count() == 5


def test_exact_commit_retry_is_idempotent(import_env):
    client, db, semester_id = import_env
    content = minimal_standard_workbook(client, semester_id)
    preview = post_workbook(client, "preview", semester_id, content).json()
    payload = {"fingerprint": preview["fingerprint"], "confirm_changes": "false"}

    first = post_workbook(client, "commit", semester_id, content, data=payload)
    retried = post_workbook(client, "commit", semester_id, content, data=payload)

    assert first.status_code == 200
    assert retried.status_code == 200
    assert retried.json()["idempotent"] is True
    assert retried.json()["batch_id"] == first.json()["batch_id"]
    assert db.query(CourseAssignment).count() == 1
    assert db.query(TeacherArrangementImportBatch).count() == 1


def test_school_codes_keep_identity_stable_when_names_change(import_env):
    client, db, semester_id = import_env
    first_content = minimal_standard_workbook(client, semester_id)
    first_preview = post_workbook(client, "preview", semester_id, first_content).json()
    first_commit = post_workbook(
        client,
        "commit",
        semester_id,
        first_content,
        data={"fingerprint": first_preview["fingerprint"], "confirm_changes": "false"},
    )
    assert first_commit.status_code == 200
    original_ids = (
        db.query(Subject).one().id,
        db.query(Teacher).one().id,
        db.query(ClassUnit).one().id,
        db.query(CourseAssignment).one().id,
    )
    changed_content = minimal_standard_workbook(client, semester_id, suffix="（新）")

    changed_preview = post_workbook(client, "preview", semester_id, changed_content)

    assert changed_preview.status_code == 200
    preview = changed_preview.json()
    assert preview["counts"]["changed"] == 3
    assert preview["counts"]["unchanged"] == 2
    rejected = post_workbook(
        client,
        "commit",
        semester_id,
        changed_content,
        data={"fingerprint": preview["fingerprint"], "confirm_changes": "false"},
    )
    assert rejected.status_code == 409
    committed = post_workbook(
        client,
        "commit",
        semester_id,
        changed_content,
        data={"fingerprint": preview["fingerprint"], "confirm_changes": "true"},
    )
    assert committed.status_code == 200, committed.json()
    assert (
        db.query(Subject).one().id,
        db.query(Teacher).one().id,
        db.query(ClassUnit).one().id,
        db.query(CourseAssignment).one().id,
    ) == original_ids
    assert db.query(Subject).one().name == "数学（新）"
    assert db.query(Teacher).one().name == "王老师（新）"
    assert db.query(ClassUnit).one().name == "七年级1班（新）"


def test_known_aliases_are_accepted_but_semester_guard_blocks_commit(import_env):
    client, db, semester_id = import_env
    content = minimal_standard_workbook(client, semester_id)
    workbook = load_workbook(io.BytesIO(content))
    workbook["科目"]["B1"] = "课程名称"
    workbook["学期"]["A4"] = 2025
    output = io.BytesIO()
    workbook.save(output)

    response = post_workbook(client, "preview", semester_id, output.getvalue())

    assert response.status_code == 200, response.json()
    preview = response.json()
    assert preview["can_commit"] is False
    issue = next(item for item in preview["issues"] if item["code"] == "semester_guard_mismatch")
    assert issue["sheet"] == "学期"
    assert issue["row"] == 4
    assert issue["field"] == "学年起始年"
    rejected = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={"fingerprint": preview["fingerprint"], "confirm_changes": "false"},
    )
    assert rejected.status_code == 409
    assert db.query(Subject).count() == 0


def test_strict_schema_and_stale_database_snapshot_are_rejected(import_env):
    client, db, semester_id = import_env
    content = minimal_standard_workbook(client, semester_id)
    workbook = load_workbook(io.BytesIO(content))
    workbook["_schema"]["B2"] = "9.9"
    output = io.BytesIO()
    workbook.save(output)

    incompatible = post_workbook(client, "preview", semester_id, output.getvalue())

    assert incompatible.status_code == 400
    assert incompatible.json()["detail"]["code"] == "template_version_unsupported"

    preview = post_workbook(client, "preview", semester_id, content).json()
    db.add(Subject(semester_id=semester_id, name="物理"))
    db.commit()
    stale = post_workbook(
        client,
        "commit",
        semester_id,
        content,
        data={"fingerprint": preview["fingerprint"], "confirm_changes": "false"},
    )
    assert stale.status_code == 409
    assert stale.json()["detail"]["code"] == "teacher_arrangement_preview_stale"


def test_database_failure_rolls_back_the_whole_import(import_env, monkeypatch):
    client, db, semester_id = import_env
    content = minimal_standard_workbook(client, semester_id)
    preview = post_workbook(client, "preview", semester_id, content).json()

    def fail_assignments(*_args, **_kwargs):
        raise IntegrityError("forced", {}, RuntimeError("forced"))

    monkeypatch.setattr(teacher_arrangement_import, "_apply_assignments", fail_assignments)
    response = post_workbook(
        client,
        "commit",
        semester_id,
        content,
        data={"fingerprint": preview["fingerprint"], "confirm_changes": "false"},
    )

    assert response.status_code == 409
    db.expire_all()
    assert db.query(Subject).count() == 0
    assert db.query(Teacher).count() == 0
    assert db.query(ClassUnit).count() == 0
    assert db.query(CourseAssignment).count() == 0


def test_standard_missing_values_are_warnings_and_require_confirmation(import_env):
    client, db, semester_id = import_env
    content = minimal_standard_workbook(client, semester_id)
    workbook = load_workbook(io.BytesIO(content))
    set_row(
        workbook,
        "教师",
        {"基础周课时": None, "行政减课时": None},
    )
    set_row(workbook, "教学任务", {"主讲教师": None})
    output = io.BytesIO()
    workbook.save(output)

    response = post_workbook(client, "preview", semester_id, output.getvalue())

    assert response.status_code == 200, response.json()
    preview = response.json()
    assert preview["can_commit"] is True
    assert preview["counts"]["warning"] == 3
    assert {issue["code"] for issue in preview["issues"]} == {
        "teacher_periods_defaulted",
        "teacher_admin_reduction_defaulted",
        "assignment_teacher_missing",
    }
    rejected = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={
            "fingerprint": preview["fingerprint"],
            "confirm_changes": "false",
            "confirm_warnings": "false",
        },
    )
    assert rejected.status_code == 409
    assert rejected.json()["detail"]["code"] == "teacher_arrangement_warnings_unconfirmed"
    committed = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={
            "fingerprint": preview["fingerprint"],
            "confirm_changes": "false",
            "confirm_warnings": "true",
        },
    )
    assert committed.status_code == 200, committed.json()
    assert db.query(Teacher).one().base_periods == 0
    assert db.query(Teacher).one().admin_reduction == 0
    assert db.query(CourseAssignment).one().teachers == []


def test_class_planned_periods_warn_in_standard_and_block_ready_mode(import_env):
    client, db, semester_id = import_env
    standard = load_workbook(
        io.BytesIO(minimal_standard_workbook(client, semester_id))
    )
    set_row(standard, "班级", {"班级计划周课时": None})
    standard_output = io.BytesIO()
    standard.save(standard_output)

    standard_preview = post_workbook(
        client, "preview", semester_id, standard_output.getvalue()
    ).json()

    assert standard_preview["can_commit"] is True
    assert standard_preview["counts"]["warning"] == 1
    assert {issue["code"] for issue in standard_preview["issues"]} == {
        "class_planned_periods_missing"
    }
    committed = post_workbook(
        client,
        "commit",
        semester_id,
        standard_output.getvalue(),
        data={
            "fingerprint": standard_preview["fingerprint"],
            "confirm_warnings": "true",
        },
    )
    assert committed.status_code == 200, committed.json()
    assert db.query(ClassUnit).one().planned_weekly_periods is None

    ready = load_workbook(
        io.BytesIO(
            minimal_standard_workbook(
                client, semester_id, suffix="-ready", mode="scheduling_ready"
            )
        )
    )
    set_row(ready, "班级", {"班级计划周课时": None})
    ready_output = io.BytesIO()
    ready.save(ready_output)

    ready_preview = post_workbook(
        client,
        "preview",
        semester_id,
        ready_output.getvalue(),
        mode="scheduling_ready",
    ).json()

    assert ready_preview["can_commit"] is False
    assert "required_value_missing" in {
        issue["code"] for issue in ready_preview["issues"]
    }


def test_ready_mode_requires_a_teacher_for_every_assignment(import_env):
    client, _, semester_id = import_env
    content = minimal_standard_workbook(
        client, semester_id, mode="scheduling_ready"
    )
    workbook = load_workbook(io.BytesIO(content))
    set_row(workbook, "教学任务", {"主讲教师": None})
    output = io.BytesIO()
    workbook.save(output)

    response = post_workbook(
        client,
        "preview",
        semester_id,
        output.getvalue(),
        mode="scheduling_ready",
    )

    assert response.status_code == 200, response.json()
    preview = response.json()
    assert preview["can_commit"] is False
    issue = next(
        item for item in preview["issues"] if item["code"] == "assignment_teacher_missing"
    )
    assert issue["severity"] == "blocker"
    assert issue["sheet"] == "教学任务"
    assert issue["row"] == 4
    assert issue["field"] == "主讲教师"


def test_compound_periods_and_ambiguous_teacher_have_precise_locations(import_env):
    client, _, semester_id = import_env
    content = minimal_standard_workbook(client, semester_id)
    workbook = load_workbook(io.BytesIO(content))
    set_row(
        workbook,
        "教师",
        {
            "学校教师编码": "T-002",
            "教师姓名": "王老师",
            "基础周课时": 18,
            "行政职务": "",
            "行政减课时": 0,
            "教师状态": "在岗",
            "外聘": "否",
        },
        row=5,
    )
    set_row(
        workbook,
        "教学任务",
        {"周课时": "4+1", "主讲教师": "王老师"},
    )
    output = io.BytesIO()
    workbook.save(output)

    response = post_workbook(client, "preview", semester_id, output.getvalue())

    assert response.status_code == 200
    issues = response.json()["issues"]
    compound = next(item for item in issues if item["code"] == "compound_periods_not_split")
    ambiguous = next(item for item in issues if item["code"] == "reference_ambiguous")
    assert compound == {
        "code": "compound_periods_not_split",
        "severity": "blocker",
        "sheet": "教学任务",
        "row": 4,
        "field": "周课时",
        "value": "4+1",
        "message": "复合课时不能写在同一任务行",
        "suggestion": "拆成多行任务，并为每行填写唯一任务编码",
    }
    assert ambiguous["sheet"] == "教学任务"
    assert ambiguous["row"] == 4
    assert ambiguous["field"] == "主讲教师"


def test_reimport_merges_workbook_changes_without_overwriting_local_changes(import_env):
    client, db, semester_id = import_env
    original = minimal_standard_workbook(client, semester_id)
    first_preview = post_workbook(client, "preview", semester_id, original).json()
    first_commit = post_workbook(
        client,
        "commit",
        semester_id,
        original,
        data={"fingerprint": first_preview["fingerprint"]},
    )
    assert first_commit.status_code == 200, first_commit.json()

    subject = db.query(Subject).one()
    subject.domain = "校内人工调整"
    db.commit()
    workbook = load_workbook(io.BytesIO(original))
    set_row(workbook, "教师", {"基础周课时": 20})
    output = io.BytesIO()
    workbook.save(output)

    response = post_workbook(client, "preview", semester_id, output.getvalue())

    assert response.status_code == 200, response.json()
    preview = response.json()
    assert preview["counts"]["changed"] == 1
    assert preview["counts"]["conflict"] == 0
    subject_row = next(
        row
        for sheet in preview["sheets"]
        if sheet["key"] == "subjects"
        for row in sheet["rows"]
    )
    assert subject_row["status"] == "unchanged"
    assert next(
        change for change in subject_row["changes"] if change["field"] == "domain"
    )["resolution"] == "system"
    committed = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={
            "fingerprint": preview["fingerprint"],
            "confirm_changes": "true",
        },
    )
    assert committed.status_code == 200, committed.json()
    db.expire_all()
    assert db.query(Subject).one().domain == "校内人工调整"
    assert db.query(Teacher).one().base_periods == 20


def test_divergent_reimport_requires_and_applies_conflict_decision(import_env):
    client, db, semester_id = import_env
    original = minimal_standard_workbook(client, semester_id)
    first_preview = post_workbook(client, "preview", semester_id, original).json()
    assert post_workbook(
        client,
        "commit",
        semester_id,
        original,
        data={"fingerprint": first_preview["fingerprint"]},
    ).status_code == 200

    teacher = db.query(Teacher).one()
    teacher.base_periods = 19
    db.commit()
    workbook = load_workbook(io.BytesIO(original))
    set_row(workbook, "教师", {"基础周课时": 20})
    output = io.BytesIO()
    workbook.save(output)

    response = post_workbook(client, "preview", semester_id, output.getvalue())

    assert response.status_code == 200, response.json()
    preview = response.json()
    assert preview["can_commit"] is False
    assert preview["counts"]["conflict"] == 1
    conflict = next(
        row
        for sheet in preview["sheets"]
        for row in sheet["rows"]
        if row["status"] == "conflict"
    )
    assert conflict["decision"] == {
        "key": "teachers:code:T-001",
        "kind": "conflict",
        "selected": None,
        "options": ["incoming", "current"],
        "removal_allowed": False,
        "reason": None,
    }
    change = next(item for item in conflict["changes"] if item["field"] == "base_periods")
    assert (change["baseline"], change["current"], change["incoming"]) == (18, 19, 20)
    rejected = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={"fingerprint": preview["fingerprint"]},
    )
    assert rejected.status_code == 409
    assert rejected.json()["detail"]["code"] == "teacher_arrangement_decisions_required"

    decisions = {conflict["decision"]["key"]: "incoming"}
    decided_preview = post_workbook(
        client,
        "preview",
        semester_id,
        output.getvalue(),
        data={"decisions": json.dumps(decisions)},
    ).json()
    assert decided_preview["fingerprint"] != preview["fingerprint"]
    assert decided_preview["can_commit"] is True

    stale = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={
            "fingerprint": preview["fingerprint"],
            "decisions": json.dumps(decisions),
            "confirm_changes": "true",
        },
    )
    assert stale.status_code == 409
    assert stale.json()["detail"]["code"] == "teacher_arrangement_preview_stale"

    committed = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={
            "fingerprint": decided_preview["fingerprint"],
            "decisions": json.dumps(decisions),
            "confirm_changes": "true",
        },
    )
    assert committed.status_code == 200, committed.json()
    db.expire_all()
    assert db.query(Teacher).one().base_periods == 20
    batch = db.query(TeacherArrangementImportBatch).order_by(
        TeacherArrangementImportBatch.id.desc()
    ).first()
    assert batch is not None and batch.decisions == decisions


def test_disappeared_source_record_requires_decision_and_can_be_removed(import_env):
    client, db, semester_id = import_env
    original = minimal_standard_workbook(client, semester_id)
    first_preview = post_workbook(client, "preview", semester_id, original).json()
    assert post_workbook(
        client,
        "commit",
        semester_id,
        original,
        data={"fingerprint": first_preview["fingerprint"]},
    ).status_code == 200
    original_target_id = db.query(TeacherArrangementSourceRecord).one().id
    workbook = load_workbook(io.BytesIO(original))
    clear_data_row(workbook, "来源记录")
    output = io.BytesIO()
    workbook.save(output)

    preview = post_workbook(client, "preview", semester_id, output.getvalue()).json()

    assert preview["counts"]["disappeared"] == 1
    disappeared = next(
        row
        for sheet in preview["sheets"]
        for row in sheet["rows"]
        if row["status"] == "disappeared"
    )
    assert disappeared["decision"]["kind"] == "disappeared"
    assert disappeared["decision"]["options"] == ["keep", "remove"]
    assert disappeared["decision"]["removal_allowed"] is True
    rejected = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={"fingerprint": preview["fingerprint"]},
    )
    assert rejected.status_code == 409
    assert rejected.json()["detail"]["code"] == "teacher_arrangement_decisions_required"

    decisions = {disappeared["decision"]["key"]: "remove"}
    decided_preview = post_workbook(
        client,
        "preview",
        semester_id,
        output.getvalue(),
        data={"decisions": json.dumps(decisions)},
    ).json()
    committed = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={
            "fingerprint": decided_preview["fingerprint"],
            "decisions": json.dumps(decisions),
        },
    )
    assert committed.status_code == 200, committed.json()
    assert committed.json()["removed"]["source_records"] == 1
    assert db.query(TeacherArrangementSourceRecord).count() == 0
    latest_record = db.query(TeacherArrangementImportRecord).order_by(
        TeacherArrangementImportRecord.id.desc()
    ).first()
    assert latest_record is not None
    assert latest_record.target_id == original_target_id
    assert latest_record.outcome == "removed"

    after = post_workbook(client, "preview", semester_id, output.getvalue()).json()
    assert after["counts"]["disappeared"] == 0


def test_disappeared_keep_detaches_source_without_deleting_target(import_env):
    client, db, semester_id = import_env
    original = minimal_standard_workbook(client, semester_id)
    first_preview = post_workbook(client, "preview", semester_id, original).json()
    assert post_workbook(
        client,
        "commit",
        semester_id,
        original,
        data={"fingerprint": first_preview["fingerprint"]},
    ).status_code == 200
    workbook = load_workbook(io.BytesIO(original))
    clear_data_row(workbook, "来源记录")
    output = io.BytesIO()
    workbook.save(output)
    preview = post_workbook(client, "preview", semester_id, output.getvalue()).json()
    disappeared = next(
        row
        for sheet in preview["sheets"]
        for row in sheet["rows"]
        if row["status"] == "disappeared"
    )

    decisions = {disappeared["decision"]["key"]: "keep"}
    decided_preview = post_workbook(
        client,
        "preview",
        semester_id,
        output.getvalue(),
        data={"decisions": json.dumps(decisions)},
    ).json()
    committed = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={
            "fingerprint": decided_preview["fingerprint"],
            "decisions": json.dumps(decisions),
        },
    )

    assert committed.status_code == 200, committed.json()
    assert committed.json()["kept"]["source_records"] == 1
    assert db.query(TeacherArrangementSourceRecord).count() == 1
    latest_record = db.query(TeacherArrangementImportRecord).order_by(
        TeacherArrangementImportRecord.id.desc()
    ).first()
    assert latest_record is not None and latest_record.outcome == "kept"
    after = post_workbook(client, "preview", semester_id, output.getvalue()).json()
    assert after["counts"]["disappeared"] == 0


def test_disappeared_referenced_teacher_cannot_be_removed(import_env):
    client, db, semester_id = import_env
    original = minimal_standard_workbook(client, semester_id)
    first_preview = post_workbook(client, "preview", semester_id, original).json()
    assert post_workbook(
        client,
        "commit",
        semester_id,
        original,
        data={"fingerprint": first_preview["fingerprint"]},
    ).status_code == 200
    workbook = load_workbook(io.BytesIO(original))
    clear_data_row(workbook, "教师")
    output = io.BytesIO()
    workbook.save(output)

    preview = post_workbook(client, "preview", semester_id, output.getvalue()).json()

    disappeared = next(
        row
        for sheet in preview["sheets"]
        for row in sheet["rows"]
        if row["status"] == "disappeared"
    )
    assert disappeared["source_key"] == "code:T-001"
    assert disappeared["decision"]["removal_allowed"] is False
    assert "引用" in disappeared["decision"]["reason"]
    decisions = {disappeared["decision"]["key"]: "remove"}
    decided_preview = post_workbook(
        client,
        "preview",
        semester_id,
        output.getvalue(),
        data={"decisions": json.dumps(decisions)},
    ).json()
    rejected = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={
            "fingerprint": decided_preview["fingerprint"],
            "decisions": json.dumps(decisions),
        },
    )
    assert rejected.status_code == 409
    assert rejected.json()["detail"]["code"] == "teacher_arrangement_blockers"
    assert db.query(Teacher).count() == 1


def test_ready_import_builds_room_and_period_data_then_requires_director_confirmation(
    import_env,
):
    client, db, semester_id = import_env
    content = complete_ready_workbook(client, semester_id)

    preview_response = post_workbook(
        client,
        "preview",
        semester_id,
        content,
        mode="scheduling_ready",
    )

    assert preview_response.status_code == 200, preview_response.json()
    preview = preview_response.json()
    assert preview["can_commit"] is True
    assert preview["counts"] == {
        "new": 7,
        "changed": 0,
        "unchanged": 0,
        "conflict": 0,
        "disappeared": 0,
        "blocker": 0,
        "warning": 0,
    }
    assert [sheet["key"] for sheet in preview["sheets"]] == [
        "subjects",
        "teachers",
        "rooms",
        "period_tables",
        "classes",
        "assignments",
        "source_records",
    ]

    committed = post_workbook(
        client,
        "commit",
        semester_id,
        content,
        data={"fingerprint": preview["fingerprint"]},
        mode="scheduling_ready",
    )

    assert committed.status_code == 200, committed.json()
    assert committed.json()["created"] == {
        "subjects": 1,
        "teachers": 1,
        "rooms": 1,
        "period_tables": 1,
        "classes": 1,
        "assignments": 1,
        "source_records": 1,
    }
    room = db.query(Room).one()
    table = db.query(PeriodTable).one()
    class_unit = db.query(ClassUnit).one()
    assert room.school_code == "ROOM-MATH"
    assert [subject.school_code for subject in room.subjects] == ["SUB-MATH"]
    assert table.school_code == "PT-JUNIOR"
    assert table.name == "初中部作息"
    assert table.num_weekdays == 5
    assert class_unit.period_table_id == table.id
    periods = db.query(Period).order_by(Period.weekday, Period.period_no).all()
    assert [(period.weekday, period.period_no) for period in periods] == [
        (1, 1),
        (2, 1),
        (3, 1),
        (4, 1),
        (5, 1),
    ]
    assert all(period.type == PeriodType.regular.value for period in periods)
    assert all(
        (period.start_time, period.end_time) == (time(8, 0), time(8, 45))
        for period in periods
    )
    assert db.get(Semester, semester_id).readiness == "draft"
    assert db.query(Timetable).count() == 0

    report = client.get(f"/api/semesters/{semester_id}/readiness")
    assert report.status_code == 200, report.json()
    assert report.json()["ready"] is False
    assert report.json()["issues"] == []
    assert [
        (check["key"], check["ok"])
        for check in report.json()["checks"]
    ] == [("data_integrity", True), ("solver_preflight", True)]

    make_user(db, "arrangement-teacher", PW, roles=[Role.teacher])
    assert client.post(
        "/api/auth/login",
        json={"username": "arrangement-teacher", "password": PW},
    ).status_code == 200
    assert client.post(f"/api/semesters/{semester_id}/readiness").status_code == 403
    assert client.post(
        "/api/auth/login",
        json={"username": "arrangement-director", "password": PW},
    ).status_code == 200

    confirmed = client.post(f"/api/semesters/{semester_id}/readiness")

    assert confirmed.status_code == 200, confirmed.json()
    assert confirmed.json()["ready"] is True
    assert db.query(Timetable).count() == 0
    audit = db.query(AuditLog).filter_by(action="confirm_semester_readiness").one()
    assert audit.username == "arrangement-director"
    assert audit.target_id == semester_id
    assert "数据完整性" in audit.detail
    assert "求解预检" in audit.detail


def test_ready_preview_blocks_period_balance_capacity_and_special_room_gaps(import_env):
    client, _, semester_id = import_env

    def issue_codes(content: bytes) -> set[str]:
        response = post_workbook(
            client,
            "preview",
            semester_id,
            content,
            mode="scheduling_ready",
        )
        assert response.status_code == 200, response.json()
        assert response.json()["can_commit"] is False
        return {issue["code"] for issue in response.json()["issues"]}

    missing_time = load_workbook(io.BytesIO(complete_ready_workbook(client, semester_id)))
    set_row(missing_time, "作息时间表", {"结束时间": None})
    missing_time_output = io.BytesIO()
    missing_time.save(missing_time_output)
    assert "required_value_missing" in issue_codes(missing_time_output.getvalue())

    mismatch = load_workbook(io.BytesIO(complete_ready_workbook(client, semester_id)))
    set_row(mismatch, "班级", {"班级计划周课时": 6})
    mismatch_output = io.BytesIO()
    mismatch.save(mismatch_output)
    assert "class_assignment_periods_mismatch" in issue_codes(mismatch_output.getvalue())

    capacity = load_workbook(io.BytesIO(complete_ready_workbook(client, semester_id)))
    set_row(capacity, "班级", {"班级计划周课时": 6})
    set_row(capacity, "教学任务", {"周课时": 6})
    capacity_output = io.BytesIO()
    capacity.save(capacity_output)
    assert "class_period_capacity_exceeded" in issue_codes(capacity_output.getvalue())

    no_special_room = load_workbook(
        io.BytesIO(complete_ready_workbook(client, semester_id))
    )
    clear_data_row(no_special_room, "教室及户外场地")
    no_special_room_output = io.BytesIO()
    no_special_room.save(no_special_room_output)
    assert "special_room_candidate_missing" in issue_codes(
        no_special_room_output.getvalue()
    )


def test_ready_reimport_replaces_period_cells_without_duplicates(import_env):
    client, db, semester_id = import_env
    original = complete_ready_workbook(client, semester_id)
    first_preview = post_workbook(
        client,
        "preview",
        semester_id,
        original,
        mode="scheduling_ready",
    ).json()
    first = post_workbook(
        client,
        "commit",
        semester_id,
        original,
        data={"fingerprint": first_preview["fingerprint"]},
        mode="scheduling_ready",
    )
    assert first.status_code == 200, first.json()
    workbook = load_workbook(io.BytesIO(original))
    set_row(workbook, "作息时间表", {"结束时间": time(8, 50)})
    output = io.BytesIO()
    workbook.save(output)

    changed_preview = post_workbook(
        client,
        "preview",
        semester_id,
        output.getvalue(),
        mode="scheduling_ready",
    )

    assert changed_preview.status_code == 200, changed_preview.json()
    assert changed_preview.json()["counts"]["changed"] == 1
    committed = post_workbook(
        client,
        "commit",
        semester_id,
        output.getvalue(),
        data={
            "fingerprint": changed_preview.json()["fingerprint"],
            "confirm_changes": "true",
        },
        mode="scheduling_ready",
    )
    assert committed.status_code == 200, committed.json()
    assert db.query(Period).count() == 5
    monday = db.query(Period).filter_by(weekday=1, period_no=1).one()
    assert monday.end_time == time(8, 50)
