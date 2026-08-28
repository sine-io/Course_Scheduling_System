"""教师安排模板下载 API 的公开契约测试。"""

import io

import pytest
from openpyxl import load_workbook

from app.api.imports import XLSX_MIME
from app.models.user import Role
from tests.conftest import make_user

PW = "password123"

STANDARD_SHEETS = ["说明", "学期", "科目", "教师", "班级", "教学任务", "来源记录"]
READY_ONLY_SHEETS = ["教室及户外场地", "作息时间表"]
DATA_SHEETS = STANDARD_SHEETS[1:]


@pytest.fixture
def template_env(env):
    client, db = env
    make_user(db, "template-director", PW, roles=[Role.director])
    client.post(
        "/api/auth/login",
        json={"username": "template-director", "password": PW},
    )
    response = client.post(
        "/api/semesters",
        json={
            "academic_year": 2026,
            "term": 1,
            "start_date": "2026-09-01",
            "end_date": "2027-01-20",
        },
    )
    assert response.status_code == 201, response.json()
    return client, response.json()["id"]


def download_template(client, semester_id: int, mode: str):
    return client.get(
        "/api/import/teacher-arrangements/template",
        params={"semester_id": semester_id, "mode": mode},
    )


@pytest.mark.parametrize(
    ("mode", "visible_sheets"),
    [
        ("standard", STANDARD_SHEETS),
        ("scheduling_ready", [*STANDARD_SHEETS, *READY_ONLY_SHEETS]),
    ],
)
def test_downloads_a_versioned_workbook_for_the_selected_semester(
    template_env, mode, visible_sheets
):
    client, semester_id = template_env

    response = download_template(client, semester_id, mode)

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MIME
    assert "teacher_arrangement_2026-2027_term-1" in response.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    assert [
        sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"
    ] == visible_sheets
    assert workbook["_schema"].sheet_state == "veryHidden"
    assert workbook["_schema"]["B1"].value == "teacher_arrangement"
    assert workbook["_schema"]["B2"].value == "1.0"
    assert workbook["_schema"]["B3"].value == mode
    assert workbook["_schema"]["B4"].value == semester_id
    schema_required = {
        (row[0], row[2]): row[5]
        for row in workbook["_schema"].iter_rows(
            min_row=10, max_col=6, values_only=True
        )
    }
    assert schema_required[("classes", "planned_weekly_periods")] is (
        mode == "scheduling_ready"
    )

    expected_data_sheets = (
        [*DATA_SHEETS, *READY_ONLY_SHEETS]
        if mode == "scheduling_ready"
        else DATA_SHEETS
    )
    for sheet_name in expected_data_sheets:
        sheet = workbook[sheet_name]
        assert all(cell.value for cell in sheet[1]), sheet_name
        assert all(cell.value for cell in sheet[2]), sheet_name
        assert any(cell.value is not None for cell in sheet[3]), sheet_name
        assert sheet.freeze_panes == "A4"
        assert sheet.auto_filter.ref is not None

    semester_sheet = workbook["学期"]
    assert [cell.value for cell in semester_sheet[1]] == ["学年起始年", "学期"]
    assert [cell.value for cell in semester_sheet[4]] == [2026, 1]

    validated_sheets = {
        sheet.title for sheet in workbook.worksheets if sheet.data_validations.count > 0
    }
    assert {"学期", "科目", "教师", "班级", "教学任务"} <= validated_sheets
    if mode == "scheduling_ready":
        assert {"教室及户外场地", "作息时间表"} <= validated_sheets


def test_template_download_rejects_unknown_semester_and_mode(template_env):
    client, semester_id = template_env

    missing = download_template(client, semester_id + 999, "standard")
    invalid_mode = download_template(client, semester_id, "anything")

    assert missing.status_code == 404
    assert missing.json()["detail"] == {
        "code": "semester_not_found",
        "message": "找不到学期",
    }
    assert invalid_mode.status_code == 422
