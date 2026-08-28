"""标准化教师安排工作簿的预览与原子提交。"""

from __future__ import annotations

import hashlib
import io
import json
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, time, timedelta
from typing import Any, Literal

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.db import Base
from app.models.assignment import AssignmentTeacher, CourseAssignment
from app.models.basedata import ClassUnit, Room, RoomType, Subject, Teacher
from app.models.period import Period, PeriodTable, PeriodType
from app.models.semester import Semester, SemesterReadiness
from app.models.teacher_arrangement_import import (
    TeacherArrangementImportBatch,
    TeacherArrangementImportRecord,
    TeacherArrangementSourceRecord,
)
from app.services.assignments import get_or_create_single_unit
from app.services.importer import ROOM_TYPE_BY_LABEL, TRACK_BY_LABEL
from app.services.teacher_arrangement_template import (
    TEMPLATE_TYPE,
    TEMPLATE_VERSION,
    FieldDefinition,
    SheetDefinition,
    TemplateMode,
    field_required,
    fields_for_mode,
    visible_definitions,
)

Severity = Literal["blocker", "warning"]
RowStatus = Literal["new", "changed", "unchanged", "conflict", "disappeared"]
DecisionKind = Literal["conflict", "disappeared"]

TEACHER_STATUS_BY_LABEL = {
    "在岗": "normal",
    "外出": "outbound",
    "产假": "maternity_leave",
    "停用": "inactive",
}
ENTITY_LABELS = {
    "subjects": "科目",
    "teachers": "教师",
    "rooms": "教室及户外场地",
    "period_tables": "作息时间表",
    "classes": "班级",
    "assignments": "教学任务",
    "source_records": "来源记录",
}
BASE_ENTITY_KEYS = (
    "subjects",
    "teachers",
    "classes",
    "assignments",
    "source_records",
)
READY_ENTITY_KEYS = (
    "subjects",
    "teachers",
    "rooms",
    "period_tables",
    "classes",
    "assignments",
    "source_records",
)
WEEKDAY_BY_LABEL = {
    "星期一": 1,
    "星期二": 2,
    "星期三": 3,
    "星期四": 4,
    "星期五": 5,
    "星期六": 6,
    "星期日": 7,
}
PERIOD_TYPE_BY_LABEL = {
    "常规课时": PeriodType.regular.value,
    "晨会": PeriodType.morning.value,
    "午休": PeriodType.lunch.value,
    "课间": PeriodType.reserved.value,
    "活动": PeriodType.reserved.value,
}


def _entity_keys(mode: TemplateMode) -> tuple[str, ...]:
    return READY_ENTITY_KEYS if mode == "scheduling_ready" else BASE_ENTITY_KEYS


class InvalidTeacherArrangementWorkbook(ValueError):
    """工作簿结构无法按已发布模板解释。"""

    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(message)

    def as_detail(self) -> dict[str, str]:
        return {"code": self.code, "message": self.message}


@dataclass(slots=True)
class LocatedIssue:
    code: str
    severity: Severity
    sheet: str
    row: int
    field: str
    value: Any
    message: str
    suggestion: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "severity": self.severity,
            "sheet": self.sheet,
            "row": self.row,
            "field": self.field,
            "value": _jsonable(self.value),
            "message": self.message,
            "suggestion": self.suggestion,
        }


@dataclass(slots=True)
class RawRow:
    sheet: str
    row: int
    values: dict[str, Any]
    headers: dict[str, str]


@dataclass(slots=True)
class PlannedRow:
    entity: str
    sheet: str
    row: int
    source_key: str
    identity: str
    values: dict[str, Any]
    raw_values: dict[str, Any]
    status: RowStatus = "new"
    changes: list[dict[str, Any]] = field(default_factory=list)
    issues: list[LocatedIssue] = field(default_factory=list)
    current_values: dict[str, Any] | None = field(default=None, repr=False)
    baseline_values: dict[str, Any] | None = field(default=None, repr=False)
    fields_to_apply: set[str] = field(default_factory=set, repr=False)
    decision_kind: DecisionKind | None = None
    decision_selected: str | None = None
    removal_allowed: bool = False
    removal_reason: str | None = None
    existing: Any = field(default=None, repr=False)
    applied: Any = field(default=None, repr=False)
    prior_record: TeacherArrangementImportRecord | None = field(default=None, repr=False)
    references: dict[str, Any] = field(default_factory=dict, repr=False)

    @property
    def decision_key(self) -> str:
        return f"{self.entity}:{self.source_key}"

    @property
    def unresolved_decision(self) -> bool:
        return self.decision_kind is not None and self.decision_selected is None

    def effective_values(self) -> dict[str, Any]:
        if self.current_values is None or self.existing is None:
            return _jsonable(self.values)
        return {
            key: _jsonable(
                value if key in self.fields_to_apply else self.current_values.get(key)
            )
            for key, value in self.values.items()
        }

    def as_dict(self) -> dict[str, Any]:
        result = {
            "sheet": self.sheet,
            "row": self.row,
            "source_key": self.source_key,
            "identity": self.identity,
            "status": self.status,
            "changes": self.changes,
            "issues": [issue.as_dict() for issue in self.issues],
        }
        if self.decision_kind is not None:
            result["decision"] = {
                "key": self.decision_key,
                "kind": self.decision_kind,
                "selected": self.decision_selected,
                "options": (
                    ["incoming", "current"]
                    if self.decision_kind == "conflict"
                    else ["keep", "remove"]
                ),
                "removal_allowed": self.removal_allowed,
                "reason": self.removal_reason,
            }
        else:
            result["decision"] = None
        return result


@dataclass(slots=True)
class ImportPlan:
    semester_id: int
    mode: TemplateMode
    workbook_sha256: str
    fingerprint: str
    decisions: dict[str, str]
    rows: dict[str, list[PlannedRow]]
    issues: list[LocatedIssue]

    @property
    def can_commit(self) -> bool:
        return not any(issue.severity == "blocker" for issue in self.issues) and not any(
            row.unresolved_decision
            for entity_rows in self.rows.values()
            for row in entity_rows
        )

    @property
    def has_unresolved_decisions(self) -> bool:
        return any(
            row.unresolved_decision
            for entity_rows in self.rows.values()
            for row in entity_rows
        )

    @property
    def has_changes(self) -> bool:
        return any(
            row.status in {"new", "changed", "conflict"}
            or (row.status == "disappeared" and row.decision_selected == "remove")
            for entity_rows in self.rows.values()
            for row in entity_rows
        )

    def counts(self) -> dict[str, int]:
        counts = {
            "new": 0,
            "changed": 0,
            "unchanged": 0,
            "conflict": 0,
            "disappeared": 0,
            "blocker": 0,
            "warning": 0,
        }
        for entity_rows in self.rows.values():
            for row in entity_rows:
                counts[row.status] += 1
        for issue in self.issues:
            counts[issue.severity] += 1
        return counts

    def as_dict(self) -> dict[str, Any]:
        return {
            "fingerprint": self.fingerprint,
            "template_version": TEMPLATE_VERSION,
            "mode": self.mode,
            "semester_id": self.semester_id,
            "can_commit": self.can_commit,
            "has_changes": self.has_changes,
            "counts": self.counts(),
            "sheets": [
                {
                    "key": entity,
                    "label": ENTITY_LABELS[entity],
                    "rows": [row.as_dict() for row in self.rows[entity]],
                }
                for entity in self.rows
            ],
            "issues": [issue.as_dict() for issue in self.issues],
        }


def _jsonable(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, dict):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_jsonable(item) for item in value]
    return str(value)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value).strip()
    return result or None


def _names(value: Any) -> tuple[str, ...]:
    text = _text(value)
    if text is None:
        return ()
    normalized = text.replace(",", "、").replace("，", "、").replace(";", "、")
    return tuple(dict.fromkeys(part.strip() for part in normalized.split("、") if part.strip()))


def _issue(
    issues: list[LocatedIssue],
    *,
    code: str,
    sheet: str,
    row: int,
    field_name: str,
    value: Any,
    message: str,
    suggestion: str,
    severity: Severity = "blocker",
    planned_row: PlannedRow | None = None,
) -> LocatedIssue:
    result = LocatedIssue(
        code=code,
        severity=severity,
        sheet=sheet,
        row=row,
        field=field_name,
        value=value,
        message=message,
        suggestion=suggestion,
    )
    issues.append(result)
    if planned_row is not None:
        planned_row.issues.append(result)
    return result


def _schema_metadata(workbook) -> dict[str, Any]:
    if "_schema" not in workbook.sheetnames:
        raise InvalidTeacherArrangementWorkbook(
            "template_schema_missing", "工作簿缺少系统模板 schema，请重新下载模板"
        )
    sheet = workbook["_schema"]
    metadata: dict[str, Any] = {}
    for row in sheet.iter_rows(min_row=1, max_row=7, max_col=2, values_only=True):
        key = _text(row[0])
        if key:
            metadata[key] = row[1]
    return metadata


def _load_workbook(
    content: bytes,
    semester: Semester,
    mode: TemplateMode,
):
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise InvalidTeacherArrangementWorkbook(
            "workbook_unreadable", "上传文件不是可读取的 XLSX 工作簿"
        ) from exc
    metadata = _schema_metadata(workbook)
    if metadata.get("template_type") != TEMPLATE_TYPE:
        raise InvalidTeacherArrangementWorkbook(
            "template_type_invalid", "上传文件不是教师安排标准化模板"
        )
    if str(metadata.get("template_version")) != TEMPLATE_VERSION:
        raise InvalidTeacherArrangementWorkbook(
            "template_version_unsupported",
            f"模板版本不受支持，请重新下载 v{TEMPLATE_VERSION} 模板",
        )
    if metadata.get("mode") != mode:
        raise InvalidTeacherArrangementWorkbook(
            "template_mode_mismatch", "上传模板类型与当前选择的导入模式不一致"
        )
    if int(metadata.get("semester_id") or 0) != semester.id:
        raise InvalidTeacherArrangementWorkbook(
            "template_semester_mismatch", "模板绑定的目标学期与当前选择不一致"
        )
    return workbook


def _field_lookup(fields: tuple[FieldDefinition, ...]) -> dict[str, FieldDefinition]:
    lookup: dict[str, FieldDefinition] = {}
    for definition in fields:
        for label in (definition.header, *definition.aliases):
            lookup[label.strip()] = definition
    return lookup


def _read_sheet_rows(
    workbook,
    definition: SheetDefinition,
    mode: TemplateMode,
    issues: list[LocatedIssue],
) -> list[RawRow]:
    fields = fields_for_mode(definition, mode)
    if definition.label not in workbook.sheetnames:
        _issue(
            issues,
            code="required_sheet_missing",
            sheet=definition.label,
            row=1,
            field_name="工作表",
            value=None,
            message=f"缺少必需工作表“{definition.label}”",
            suggestion="请使用系统下载的原始模板补齐工作表",
        )
        return []
    sheet = workbook[definition.label]
    lookup = _field_lookup(fields)
    columns: dict[int, FieldDefinition] = {}
    seen_fields: set[str] = set()
    for column, cell in enumerate(sheet[1], start=1):
        label = _text(cell.value)
        if label is None:
            continue
        field_definition = lookup.get(label)
        if field_definition is None:
            _issue(
                issues,
                code="unknown_header",
                sheet=definition.label,
                row=1,
                field_name=label,
                value=label,
                message=f"模板不支持表头“{label}”",
                suggestion="请恢复系统模板表头；仅支持模板内声明的已知别名",
            )
            continue
        if field_definition.key in seen_fields:
            _issue(
                issues,
                code="duplicate_header",
                sheet=definition.label,
                row=1,
                field_name=label,
                value=label,
                message=f"字段“{field_definition.header}”出现多次",
                suggestion="每个模板字段只保留一列",
            )
            continue
        seen_fields.add(field_definition.key)
        columns[column] = field_definition
    for field_definition in fields:
        if field_required(field_definition, mode) and field_definition.key not in seen_fields:
            _issue(
                issues,
                code="required_header_missing",
                sheet=definition.label,
                row=1,
                field_name=field_definition.header,
                value=None,
                message=f"缺少必填字段“{field_definition.header}”",
                suggestion="请重新下载模板或恢复该列",
            )

    result: list[RawRow] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
        values: dict[str, Any] = {
            field_definition.key: cells[column - 1] if column <= len(cells) else None
            for column, field_definition in columns.items()
        }
        if not any(_text(value) is not None for value in values.values()):
            continue
        headers = {
            field_definition.key: field_definition.header
            for field_definition in columns.values()
        }
        raw_row = RawRow(definition.label, row_number, values, headers)
        result.append(raw_row)
        for field_definition in fields:
            if (
                field_required(field_definition, mode)
                and field_definition.key in seen_fields
                and _text(values.get(field_definition.key)) is None
            ):
                _issue(
                    issues,
                    code="required_value_missing",
                    sheet=definition.label,
                    row=row_number,
                    field_name=field_definition.header,
                    value=None,
                    message=f"“{field_definition.header}”为必填项",
                    suggestion="请填写后重新上传",
                )
    return result


def _integer(
    raw_row: RawRow,
    key: str,
    issues: list[LocatedIssue],
    *,
    default: int | None = None,
    minimum: int = 0,
    maximum: int = 10000,
) -> int | None:
    value = raw_row.values.get(key)
    text = _text(value)
    if text is None:
        return default
    if key == "weekly_periods" and "+" in text:
        _issue(
            issues,
            code="compound_periods_not_split",
            sheet=raw_row.sheet,
            row=raw_row.row,
            field_name=raw_row.headers.get(key, key),
            value=value,
            message="复合课时不能写在同一任务行",
            suggestion="拆成多行任务，并为每行填写唯一任务编码",
        )
        return default
    try:
        number = float(text)
    except ValueError:
        number = -1.5
    if not number.is_integer() or not minimum <= number <= maximum:
        _issue(
            issues,
            code="invalid_integer",
            sheet=raw_row.sheet,
            row=raw_row.row,
            field_name=raw_row.headers.get(key, key),
            value=value,
            message=f"“{text}”不是 {minimum} 至 {maximum} 之间的整数",
            suggestion="请填写单个整数；复合课时须拆成多行任务",
        )
        return default
    return int(number)


def _choice(
    raw_row: RawRow,
    key: str,
    choices: dict[str, str],
    issues: list[LocatedIssue],
    *,
    default: str,
) -> str:
    value = _text(raw_row.values.get(key))
    if value is None:
        return default
    if value not in choices:
        _issue(
            issues,
            code="invalid_choice",
            sheet=raw_row.sheet,
            row=raw_row.row,
            field_name=raw_row.headers.get(key, key),
            value=value,
            message=f"“{value}”不是支持的选项",
            suggestion=f"请填写：{'、'.join(choices)}",
        )
        return default
    return choices[value]


def _boolean(raw_row: RawRow, key: str, issues: list[LocatedIssue]) -> bool:
    return _choice(
        raw_row,
        key,
        {"是": "true", "否": "false"},
        issues,
        default="false",
    ) == "true"


def _clock_time(
    raw_row: RawRow,
    key: str,
    issues: list[LocatedIssue],
) -> str | None:
    value = raw_row.values.get(key)
    if _text(value) is None:
        return None
    parsed: time | None = None
    if isinstance(value, datetime):
        parsed = value.time()
    elif isinstance(value, time):
        parsed = value
    elif isinstance(value, timedelta):
        seconds = int(value.total_seconds()) % (24 * 60 * 60)
        parsed = time(seconds // 3600, (seconds % 3600) // 60, seconds % 60)
    elif isinstance(value, (int, float)) and 0 <= value < 1:
        seconds = round(float(value) * 24 * 60 * 60) % (24 * 60 * 60)
        parsed = time(seconds // 3600, (seconds % 3600) // 60, seconds % 60)
    else:
        text = _text(value) or ""
        for fmt in ("%H:%M", "%H:%M:%S"):
            try:
                parsed = datetime.strptime(text, fmt).time()
                break
            except ValueError:
                continue
    if parsed is None:
        _issue(
            issues,
            code="invalid_time",
            sheet=raw_row.sheet,
            row=raw_row.row,
            field_name=raw_row.headers.get(key, key),
            value=value,
            message="时间格式无效",
            suggestion="请使用 Excel 时间格式，例如 08:00",
        )
        return None
    return parsed.replace(microsecond=0).isoformat()


def _time_from_iso(value: str | None) -> time | None:
    return time.fromisoformat(value) if value else None


def _source_key(code: str | None, name: str) -> str:
    return f"code:{code}" if code else f"name:{name}"


def _object_reference(target: Any) -> str:
    code = _text(getattr(target, "school_code", None))
    return code or str(target.name)


def _planned_reference(target: PlannedRow | Any) -> str:
    if isinstance(target, PlannedRow):
        return _text(target.values.get("school_code")) or str(target.values["name"])
    return _object_reference(target)


def _identify_existing(
    *,
    row: PlannedRow,
    existing: list[Any],
    issues: list[LocatedIssue],
) -> Any:
    code = _text(row.values.get("school_code"))
    name = str(row.values["name"])
    if code:
        code_matches = [item for item in existing if _text(item.school_code) == code]
        if code_matches:
            match = code_matches[0]
            name_collision = [
                item for item in existing if item.name == name and item.id != match.id
            ]
            if name_collision:
                _issue(
                    issues,
                    code="identity_name_conflict",
                    sheet=row.sheet,
                    row=row.row,
                    field_name="名称",
                    value=name,
                    message="该名称已属于另一个学校编码",
                    suggestion="请使用唯一名称，或核对学校编码",
                    planned_row=row,
                )
            return match
        name_matches = [item for item in existing if item.name == name]
        if len(name_matches) == 1 and not _text(name_matches[0].school_code):
            return name_matches[0]
        if name_matches:
            _issue(
                issues,
                code="identity_code_conflict",
                sheet=row.sheet,
                row=row.row,
                field_name="学校编码",
                value=code,
                message="该名称已绑定其他学校编码",
                suggestion="已有编码一旦建立即优先，请核对编码",
                planned_row=row,
            )
        return None
    name_matches = [item for item in existing if item.name == name]
    if len(name_matches) == 1:
        return name_matches[0]
    if len(name_matches) > 1:
        _issue(
            issues,
            code="identity_name_ambiguous",
            sheet=row.sheet,
            row=row.row,
            field_name="名称",
            value=name,
            message="目标学期内存在多个同名对象，无法唯一识别",
            suggestion="请为对象填写学校编码并在引用处使用编码",
            planned_row=row,
        )
    return None


def _set_status(row: PlannedRow, current: dict[str, Any] | None) -> None:
    row.current_values = current
    if current is None:
        row.status = "new"
        row.fields_to_apply = set(row.values)
        return
    changes = []
    for key, after in row.values.items():
        before = current.get(key)
        if _jsonable(before) != _jsonable(after):
            changes.append({"field": key, "before": _jsonable(before), "after": _jsonable(after)})
    row.changes = changes
    row.fields_to_apply = {change["field"] for change in changes}
    row.status = "changed" if changes else "unchanged"


def _deduplicate_named_rows(rows: list[PlannedRow], issues: list[LocatedIssue]) -> None:
    by_code: dict[str, list[PlannedRow]] = defaultdict(list)
    by_name_without_code: dict[str, list[PlannedRow]] = defaultdict(list)
    for row in rows:
        code = _text(row.values.get("school_code"))
        if code:
            by_code[code].append(row)
        else:
            by_name_without_code[str(row.values["name"])].append(row)
    for value, duplicates in [*by_code.items(), *by_name_without_code.items()]:
        if len(duplicates) < 2:
            continue
        for row in duplicates:
            _issue(
                issues,
                code="duplicate_identity",
                sheet=row.sheet,
                row=row.row,
                field_name="学校编码/名称",
                value=value,
                message="工作簿内身份重复",
                suggestion="每个对象只保留一行，并确保编码唯一",
                planned_row=row,
            )


def _plan_subjects(
    db: Session, semester_id: int, raw_rows: list[RawRow], issues: list[LocatedIssue]
) -> list[PlannedRow]:
    existing = list(db.scalars(select(Subject).where(Subject.semester_id == semester_id)))
    result: list[PlannedRow] = []
    room_types = {label: value.value for label, value in ROOM_TYPE_BY_LABEL.items()}
    for raw in raw_rows:
        name = _text(raw.values.get("name")) or ""
        code = _text(raw.values.get("school_code"))
        values = {
            "school_code": code,
            "name": name,
            "domain": _text(raw.values.get("domain")),
            "required_room_type": _choice(
                raw, "required_room_type", room_types, issues, default="normal"
            ),
        }
        row = PlannedRow(
            "subjects",
            raw.sheet,
            raw.row,
            _source_key(code, name),
            code or name,
            values,
            raw.values,
        )
        row.existing = _identify_existing(row=row, existing=existing, issues=issues)
        current = None
        if row.existing is not None:
            current = {
                "school_code": row.existing.school_code,
                "name": row.existing.name,
                "domain": row.existing.domain,
                "required_room_type": row.existing.required_room_type or "normal",
            }
        _set_status(row, current)
        result.append(row)
    _deduplicate_named_rows(result, issues)
    return result


def _plan_teachers(
    db: Session, semester_id: int, raw_rows: list[RawRow], issues: list[LocatedIssue]
) -> list[PlannedRow]:
    existing = list(db.scalars(select(Teacher).where(Teacher.semester_id == semester_id)))
    result: list[PlannedRow] = []
    for raw in raw_rows:
        name = _text(raw.values.get("name")) or ""
        code = _text(raw.values.get("school_code"))
        status = _choice(
            raw,
            "status",
            TEACHER_STATUS_BY_LABEL,
            issues,
            default="normal",
        )
        values = {
            "school_code": code,
            "name": name,
            "base_periods": _integer(raw, "base_periods", issues, default=0, maximum=100),
            "admin_title": _text(raw.values.get("admin_title")),
            "admin_reduction": _integer(
                raw, "admin_reduction", issues, default=0, maximum=100
            ),
            "arrangement_status": status,
            "is_external": _boolean(raw, "is_external", issues),
            "is_active": status != "inactive",
        }
        row = PlannedRow(
            "teachers",
            raw.sheet,
            raw.row,
            _source_key(code, name),
            code or name,
            values,
            raw.values,
        )
        if _text(raw.values.get("base_periods")) is None:
            _issue(
                issues,
                code="teacher_periods_defaulted",
                sheet=row.sheet,
                row=row.row,
                field_name="基础周课时",
                value=None,
                message="基础周课时未填写，将按 0 导入",
                suggestion="如需核算教师工作量，请填写实际基础周课时",
                severity="warning",
                planned_row=row,
            )
        if _text(raw.values.get("admin_reduction")) is None:
            _issue(
                issues,
                code="teacher_admin_reduction_defaulted",
                sheet=row.sheet,
                row=row.row,
                field_name="行政减课时",
                value=None,
                message="行政减课时未填写，将按 0 导入",
                suggestion="有行政减课时请填写非负整数；系统不会从备注推断",
                severity="warning",
                planned_row=row,
            )
        row.existing = _identify_existing(row=row, existing=existing, issues=issues)
        current = None
        if row.existing is not None:
            current = {
                "school_code": row.existing.school_code,
                "name": row.existing.name,
                "base_periods": row.existing.base_periods,
                "admin_title": row.existing.admin_title,
                "admin_reduction": row.existing.admin_reduction,
                "arrangement_status": row.existing.arrangement_status,
                "is_external": row.existing.is_external,
                "is_active": row.existing.is_active,
            }
        _set_status(row, current)
        result.append(row)
    _deduplicate_named_rows(result, issues)
    return result


def _reference_candidates(planned: list[PlannedRow], existing: list[Any], value: str) -> list[Any]:
    planned_by_code = [row for row in planned if _text(row.values.get("school_code")) == value]
    if planned_by_code:
        return planned_by_code
    existing_by_code = [item for item in existing if _text(item.school_code) == value]
    if existing_by_code:
        return existing_by_code
    return [row for row in planned if row.values.get("name") == value] + [
        item for item in existing if item.name == value
    ]


def _resolve_reference(
    *,
    value: str | None,
    entity_label: str,
    planned: list[PlannedRow],
    existing: list[Any],
    issues: list[LocatedIssue],
    owner: PlannedRow,
    field_name: str,
    required: bool,
) -> Any:
    if value is None:
        if required:
            _issue(
                issues,
                code="reference_missing",
                sheet=owner.sheet,
                row=owner.row,
                field_name=field_name,
                value=None,
                message=f"缺少{entity_label}引用",
                suggestion=f"填写已存在或本工作簿内的{entity_label}编码/唯一名称",
                planned_row=owner,
            )
        return None
    candidates = _reference_candidates(planned, existing, value)
    unique = []
    seen: set[tuple[str, int]] = set()
    for candidate in candidates:
        key = (
            "planned" if isinstance(candidate, PlannedRow) else "existing",
            id(candidate) if isinstance(candidate, PlannedRow) else candidate.id,
        )
        if key not in seen:
            seen.add(key)
            unique.append(candidate)
    if len(unique) == 1:
        return unique[0]
    code = "reference_not_found" if not unique else "reference_ambiguous"
    message = f"找不到{entity_label}“{value}”" if not unique else f"{entity_label}“{value}”不唯一"
    _issue(
        issues,
        code=code,
        sheet=owner.sheet,
        row=owner.row,
        field_name=field_name,
        value=value,
        message=message,
        suggestion=f"请填写{entity_label}学校编码；无编码时名称必须唯一",
        planned_row=owner,
    )
    return None


def _plan_rooms(
    db: Session,
    semester_id: int,
    raw_rows: list[RawRow],
    subject_rows: list[PlannedRow],
    issues: list[LocatedIssue],
) -> list[PlannedRow]:
    existing = list(db.scalars(select(Room).where(Room.semester_id == semester_id)))
    existing_subjects = list(
        db.scalars(select(Subject).where(Subject.semester_id == semester_id))
    )
    room_types = {label: value.value for label, value in ROOM_TYPE_BY_LABEL.items()}
    result: list[PlannedRow] = []
    for raw in raw_rows:
        name = _text(raw.values.get("name")) or ""
        code = _text(raw.values.get("school_code"))
        values: dict[str, Any] = {
            "school_code": code,
            "name": name,
            "room_type": _choice(
                raw,
                "room_type",
                room_types,
                issues,
                default=RoomType.normal.value,
            ),
            "capacity": _integer(raw, "capacity", issues, default=None, maximum=5000),
            "applicable_subjects": [],
        }
        row = PlannedRow(
            "rooms",
            raw.sheet,
            raw.row,
            _source_key(code, name),
            code or name,
            values,
            raw.values,
        )
        targets: list[Any] = []
        for subject_ref in _names(raw.values.get("applicable_subjects")):
            target = _resolve_reference(
                value=subject_ref,
                entity_label="科目",
                planned=subject_rows,
                existing=existing_subjects,
                issues=issues,
                owner=row,
                field_name="适用科目",
                required=True,
            )
            if target is not None and target not in targets:
                targets.append(target)
        row.references["applicable_subjects"] = targets
        row.values["applicable_subjects"] = sorted(
            _planned_reference(target) for target in targets
        )
        row.existing = _identify_existing(row=row, existing=existing, issues=issues)
        current = None
        if row.existing is not None:
            current = {
                "school_code": row.existing.school_code,
                "name": row.existing.name,
                "room_type": row.existing.room_type,
                "capacity": row.existing.capacity,
                "applicable_subjects": sorted(
                    _object_reference(subject) for subject in row.existing.subjects
                ),
            }
        _set_status(row, current)
        result.append(row)
    _deduplicate_named_rows(result, issues)
    return result


def _period_name(period_type: str, period_number: int) -> str:
    if period_type == PeriodType.regular.value:
        return f"第{period_number}节"
    return {
        PeriodType.morning.value: "晨会",
        PeriodType.lunch.value: "午休",
        PeriodType.homeroom.value: "班主任时间",
        PeriodType.reserved.value: "活动",
    }.get(period_type, f"第{period_number}节")


def _period_values(period: Period) -> dict[str, Any]:
    return {
        "weekday": period.weekday,
        "period_number": period.period_no,
        "period_type": period.type,
        "start_time": (
            period.start_time.replace(microsecond=0).isoformat()
            if period.start_time
            else None
        ),
        "end_time": (
            period.end_time.replace(microsecond=0).isoformat()
            if period.end_time
            else None
        ),
    }


def _period_table_current(table: PeriodTable) -> dict[str, Any]:
    return {
        "school_code": table.school_code,
        "name": table.name,
        "num_weekdays": table.num_weekdays,
        "is_default": table.is_default,
        "periods": [
            _period_values(period)
            for period in sorted(
                table.periods,
                key=lambda item: (item.weekday, item.period_no),
            )
        ],
    }


def _plan_period_tables(
    db: Session,
    semester_id: int,
    raw_rows: list[RawRow],
    issues: list[LocatedIssue],
) -> list[PlannedRow]:
    existing = list(
        db.scalars(select(PeriodTable).where(PeriodTable.semester_id == semester_id))
    )
    grouped: dict[str, list[RawRow]] = defaultdict(list)
    for raw in raw_rows:
        grouped[_text(raw.values.get("table_code")) or ""].append(raw)

    result: list[PlannedRow] = []
    has_existing_default = any(table.is_default for table in existing)
    for index, (code, group) in enumerate(grouped.items()):
        first = group[0]
        name = _text(first.values.get("table_name")) or ""
        values: dict[str, Any] = {
            "school_code": code,
            "name": name,
            "num_weekdays": 0,
            "is_default": False,
            "periods": [],
        }
        row = PlannedRow(
            "period_tables",
            first.sheet,
            first.row,
            _source_key(code, name),
            code or name,
            values,
            {"rows": [raw.values for raw in group]},
        )
        seen_cells: set[tuple[int, int]] = set()
        periods: list[dict[str, Any]] = []
        for raw in group:
            row_name = _text(raw.values.get("table_name")) or ""
            if row_name != name:
                _issue(
                    issues,
                    code="period_table_name_inconsistent",
                    sheet=raw.sheet,
                    row=raw.row,
                    field_name="作息表名称",
                    value=row_name,
                    message="同一作息表编码使用了不同名称",
                    suggestion=f"将名称统一为“{name}”",
                    planned_row=row,
                )
            weekday = int(
                _choice(
                    raw,
                    "weekday",
                    {label: str(value) for label, value in WEEKDAY_BY_LABEL.items()},
                    issues,
                    default="1",
                )
            )
            period_number = _integer(
                raw,
                "period_number",
                issues,
                minimum=1,
                maximum=20,
            )
            period_type = _choice(
                raw,
                "period_type",
                PERIOD_TYPE_BY_LABEL,
                issues,
                default=PeriodType.regular.value,
            )
            start_time = _clock_time(raw, "start_time", issues)
            end_time = _clock_time(raw, "end_time", issues)
            if start_time and end_time and end_time <= start_time:
                _issue(
                    issues,
                    code="period_time_order_invalid",
                    sheet=raw.sheet,
                    row=raw.row,
                    field_name="结束时间",
                    value=end_time,
                    message="结束时间必须晚于开始时间",
                    suggestion="核对该节次的开始和结束时间",
                    planned_row=row,
                )
            if period_number is None:
                continue
            cell = (weekday, period_number)
            if cell in seen_cells:
                _issue(
                    issues,
                    code="period_cell_duplicate",
                    sheet=raw.sheet,
                    row=raw.row,
                    field_name="星期/节次",
                    value=f"{weekday}-{period_number}",
                    message="同一作息表中的星期和节次重复",
                    suggestion="每套作息表的每个星期、节次只保留一行",
                    planned_row=row,
                )
                continue
            seen_cells.add(cell)
            periods.append(
                {
                    "weekday": weekday,
                    "period_number": period_number,
                    "period_type": period_type,
                    "start_time": start_time,
                    "end_time": end_time,
                }
            )
        periods.sort(key=lambda item: (item["weekday"], item["period_number"]))
        row.values["periods"] = periods
        row.values["num_weekdays"] = max(
            (period["weekday"] for period in periods),
            default=0,
        )
        row.existing = _identify_existing(row=row, existing=existing, issues=issues)
        row.values["is_default"] = bool(
            row.existing.is_default
            if row.existing is not None
            else not has_existing_default and index == 0
        )
        _set_status(
            row,
            _period_table_current(row.existing) if row.existing is not None else None,
        )
        result.append(row)
    _deduplicate_named_rows(result, issues)
    return result


def _plan_classes(
    db: Session,
    semester_id: int,
    raw_rows: list[RawRow],
    teacher_rows: list[PlannedRow],
    period_table_rows: list[PlannedRow],
    issues: list[LocatedIssue],
    mode: TemplateMode,
) -> list[PlannedRow]:
    existing = list(db.scalars(select(ClassUnit).where(ClassUnit.semester_id == semester_id)))
    existing_teachers = list(
        db.scalars(select(Teacher).where(Teacher.semester_id == semester_id))
    )
    existing_period_tables = list(
        db.scalars(select(PeriodTable).where(PeriodTable.semester_id == semester_id))
    )
    tracks = {label: value.value for label, value in TRACK_BY_LABEL.items()}
    result: list[PlannedRow] = []
    for raw in raw_rows:
        name = _text(raw.values.get("name")) or ""
        code = _text(raw.values.get("school_code"))
        values = {
            "school_code": code,
            "name": name,
            "grade": _integer(raw, "grade", issues, minimum=1, maximum=12),
            "track": _choice(raw, "track", tracks, issues, default="junior_high"),
            "department": _text(raw.values.get("specialization")),
            "planned_weekly_periods": _integer(
                raw, "planned_weekly_periods", issues, minimum=0, maximum=200
            ),
            "homeroom_teacher": None,
        }
        if mode == "scheduling_ready":
            values["period_table"] = None
        row = PlannedRow(
            "classes",
            raw.sheet,
            raw.row,
            _source_key(code, name),
            code or name,
            values,
            raw.values,
        )
        if (
            mode == "standard"
            and _text(raw.values.get("planned_weekly_periods")) is None
        ):
            _issue(
                issues,
                code="class_planned_periods_missing",
                sheet=row.sheet,
                row=row.row,
                field_name="班级计划周课时",
                value=None,
                message="班级计划周课时未填写",
                suggestion="自动排课前请补齐该班所有可排教学任务的周课时目标总数",
                severity="warning",
                planned_row=row,
            )
        homeroom_ref = _text(raw.values.get("homeroom_teacher"))
        target = _resolve_reference(
            value=homeroom_ref,
            entity_label="教师",
            planned=teacher_rows,
            existing=existing_teachers,
            issues=issues,
            owner=row,
            field_name="班主任",
            required=False,
        )
        row.references["homeroom_teacher"] = target
        row.values["homeroom_teacher"] = _planned_reference(target) if target else None
        if mode == "scheduling_ready":
            period_table_target = _resolve_reference(
                value=_text(raw.values.get("period_table_code")),
                entity_label="作息时间表",
                planned=period_table_rows,
                existing=existing_period_tables,
                issues=issues,
                owner=row,
                field_name="作息表编码",
                required=True,
            )
            row.references["period_table"] = period_table_target
            row.values["period_table"] = (
                _planned_reference(period_table_target)
                if period_table_target is not None
                else None
            )
        row.existing = _identify_existing(row=row, existing=existing, issues=issues)
        current = None
        if row.existing is not None:
            current = {
                "school_code": row.existing.school_code,
                "name": row.existing.name,
                "grade": row.existing.grade,
                "track": row.existing.track,
                "department": row.existing.department,
                "planned_weekly_periods": row.existing.planned_weekly_periods,
                "homeroom_teacher": (
                    _object_reference(row.existing.homeroom_teacher)
                    if row.existing.homeroom_teacher
                    else None
                ),
            }
            if mode == "scheduling_ready":
                current["period_table"] = (
                    _object_reference(row.existing.period_table)
                    if row.existing.period_table
                    else None
                )
        _set_status(row, current)
        result.append(row)
    _deduplicate_named_rows(result, issues)
    return result


def _assignment_current(existing: CourseAssignment) -> dict[str, Any]:
    members = existing.scheduling_unit.members
    class_ref = _object_reference(members[0].class_unit) if len(members) == 1 else None
    lead = next((link.teacher for link in existing.teachers if link.is_lead), None)
    co_teachers = sorted(
        _object_reference(link.teacher) for link in existing.teachers if not link.is_lead
    )
    return {
        "task_code": existing.task_code,
        "class_ref": class_ref,
        "subject_ref": _object_reference(existing.subject),
        "component": existing.component,
        "weekly_periods": existing.periods_per_week,
        "lead_teacher": _object_reference(lead) if lead else None,
        "co_teachers": co_teachers,
    }


def _plan_assignments(
    db: Session,
    semester_id: int,
    raw_rows: list[RawRow],
    subject_rows: list[PlannedRow],
    teacher_rows: list[PlannedRow],
    class_rows: list[PlannedRow],
    issues: list[LocatedIssue],
    mode: TemplateMode,
) -> list[PlannedRow]:
    existing = list(
        db.scalars(select(CourseAssignment).where(CourseAssignment.semester_id == semester_id))
    )
    existing_subjects = list(
        db.scalars(select(Subject).where(Subject.semester_id == semester_id))
    )
    existing_teachers = list(
        db.scalars(select(Teacher).where(Teacher.semester_id == semester_id))
    )
    existing_classes = list(
        db.scalars(select(ClassUnit).where(ClassUnit.semester_id == semester_id))
    )
    result: list[PlannedRow] = []
    seen_codes: dict[str, PlannedRow] = {}
    for raw in raw_rows:
        task_code = _text(raw.values.get("task_code")) or ""
        values: dict[str, Any] = {
            "task_code": task_code,
            "class_ref": None,
            "subject_ref": None,
            "component": _text(raw.values.get("component")) or "",
            "weekly_periods": _integer(
                raw, "weekly_periods", issues, minimum=1, maximum=40
            ),
            "lead_teacher": None,
            "co_teachers": [],
        }
        row = PlannedRow(
            "assignments",
            raw.sheet,
            raw.row,
            task_code,
            task_code,
            values,
            raw.values,
        )
        if task_code in seen_codes:
            _issue(
                issues,
                code="duplicate_task_code",
                sheet=row.sheet,
                row=row.row,
                field_name="任务编码",
                value=task_code,
                message="任务编码在工作簿内重复",
                suggestion="每个可独立排课任务使用唯一编码",
                planned_row=row,
            )
        else:
            seen_codes[task_code] = row

        class_target = _resolve_reference(
            value=_text(raw.values.get("class_ref")),
            entity_label="班级",
            planned=class_rows,
            existing=existing_classes,
            issues=issues,
            owner=row,
            field_name="班级",
            required=True,
        )
        subject_target = _resolve_reference(
            value=_text(raw.values.get("subject_ref")),
            entity_label="科目",
            planned=subject_rows,
            existing=existing_subjects,
            issues=issues,
            owner=row,
            field_name="科目",
            required=True,
        )
        lead_target = _resolve_reference(
            value=_text(raw.values.get("lead_teacher")),
            entity_label="教师",
            planned=teacher_rows,
            existing=existing_teachers,
            issues=issues,
            owner=row,
            field_name="主讲教师",
            required=False,
        )
        if _text(raw.values.get("lead_teacher")) is None:
            _issue(
                issues,
                code="assignment_teacher_missing",
                sheet=row.sheet,
                row=row.row,
                field_name="主讲教师",
                value=None,
                message=(
                    "自动排课准备模式下教学任务必须指定主讲教师"
                    if mode == "scheduling_ready"
                    else "教学任务尚未指定主讲教师"
                ),
                suggestion="填写教师学校编码或目标学期内唯一姓名",
                severity="blocker" if mode == "scheduling_ready" else "warning",
                planned_row=row,
            )
        co_targets = []
        for teacher_ref in _names(raw.values.get("co_teachers")):
            target = _resolve_reference(
                value=teacher_ref,
                entity_label="教师",
                planned=teacher_rows,
                existing=existing_teachers,
                issues=issues,
                owner=row,
                field_name="协同教师",
                required=True,
            )
            if target is not None and target is not lead_target:
                co_targets.append(target)
        row.references.update(
            {
                "class": class_target,
                "subject": subject_target,
                "lead_teacher": lead_target,
                "co_teachers": co_targets,
            }
        )
        row.values.update(
            {
                "class_ref": _planned_reference(class_target) if class_target else None,
                "subject_ref": _planned_reference(subject_target) if subject_target else None,
                "lead_teacher": _planned_reference(lead_target) if lead_target else None,
                "co_teachers": sorted(_planned_reference(target) for target in co_targets),
            }
        )
        matches = [assignment for assignment in existing if assignment.task_code == task_code]
        row.existing = matches[0] if matches else None
        _set_status(row, _assignment_current(row.existing) if row.existing else None)
        result.append(row)
    return result


def _regular_period_capacity(target: PlannedRow | PeriodTable | None) -> int:
    if target is None:
        return 0
    if isinstance(target, PlannedRow):
        return sum(
            1
            for period in target.values.get("periods", [])
            if period.get("period_type") == PeriodType.regular.value
        )
    return sum(1 for period in target.periods if period.type == PeriodType.regular.value)


def _room_accepts_subject(
    room: PlannedRow | Room,
    subject: PlannedRow | Subject,
    required_room_type: str,
) -> bool:
    if isinstance(room, PlannedRow):
        if room.values.get("room_type") != required_room_type:
            return False
        subjects = room.references.get("applicable_subjects", [])
        return not subjects or subject in subjects
    if room.room_type != required_room_type:
        return False
    subject_id = (
        subject.existing.id
        if isinstance(subject, PlannedRow) and subject.existing is not None
        else subject.id
        if isinstance(subject, Subject)
        else None
    )
    return not room.subjects or any(item.id == subject_id for item in room.subjects)


def _validate_ready_plan(
    db: Session,
    semester_id: int,
    rows: dict[str, list[PlannedRow]],
    issues: list[LocatedIssue],
) -> None:
    assignment_rows = rows["assignments"]
    for class_row in rows["classes"]:
        planned = class_row.values.get("planned_weekly_periods")
        assigned = sum(
            assignment.values.get("weekly_periods") or 0
            for assignment in assignment_rows
            if assignment.references.get("class") is class_row
        )
        if planned is not None and assigned != planned:
            _issue(
                issues,
                code="class_assignment_periods_mismatch",
                sheet=class_row.sheet,
                row=class_row.row,
                field_name="班级计划周课时",
                value={"planned": planned, "assigned": assigned},
                message=f"班级计划周课时为 {planned}，教学任务合计为 {assigned}",
                suggestion="调整教学任务周课时，使合计与班级计划完全一致",
                planned_row=class_row,
            )
        capacity = _regular_period_capacity(class_row.references.get("period_table"))
        if planned is not None and planned > capacity:
            _issue(
                issues,
                code="class_period_capacity_exceeded",
                sheet=class_row.sheet,
                row=class_row.row,
                field_name="班级计划周课时",
                value={"planned": planned, "capacity": capacity},
                message=f"班级计划周课时 {planned} 超过作息表常规课时容量 {capacity}",
                suggestion="增加作息表常规课时，或降低班级计划周课时",
                planned_row=class_row,
            )

    planned_rooms = [
        room
        for room in rows["rooms"]
        if room.status != "disappeared" or room.decision_selected == "keep"
    ]
    replaced_room_ids = {
        room.existing.id for room in planned_rooms if room.existing is not None
    }
    existing_rooms = [
        room
        for room in db.scalars(select(Room).where(Room.semester_id == semester_id))
        if room.id not in replaced_room_ids
    ]
    room_candidates: list[PlannedRow | Room] = [*planned_rooms, *existing_rooms]
    for assignment in assignment_rows:
        subject = assignment.references.get("subject")
        if subject is None:
            continue
        required_room_type = (
            subject.values.get("required_room_type")
            if isinstance(subject, PlannedRow)
            else subject.required_room_type
        ) or RoomType.normal.value
        if required_room_type == RoomType.normal.value:
            continue
        if any(
            _room_accepts_subject(room, subject, required_room_type)
            for room in room_candidates
        ):
            continue
        _issue(
            issues,
            code="special_room_candidate_missing",
            sheet=assignment.sheet,
            row=assignment.row,
            field_name="科目",
            value=assignment.values.get("subject_ref"),
            message="该教学任务需要专用教室或户外场地，但没有适用的候选教室资源",
            suggestion="在“教室及户外场地”中补充同类型且适用于该科目的教室资源",
            planned_row=assignment,
        )


def _plan_source_records(
    db: Session, semester_id: int, raw_rows: list[RawRow], issues: list[LocatedIssue]
) -> list[PlannedRow]:
    existing = list(
        db.scalars(
            select(TeacherArrangementSourceRecord).where(
                TeacherArrangementSourceRecord.semester_id == semester_id
            )
        )
    )
    result: list[PlannedRow] = []
    seen_codes: set[str] = set()
    for raw in raw_rows:
        code = _text(raw.values.get("record_code")) or ""
        values = {
            "record_code": code,
            "category": _text(raw.values.get("category")) or "",
            "task_code": _text(raw.values.get("task_code")),
            "content": _text(raw.values.get("content")) or "",
            "notes": _text(raw.values.get("notes")),
        }
        row = PlannedRow(
            "source_records", raw.sheet, raw.row, code, code, values, raw.values
        )
        if code in seen_codes:
            _issue(
                issues,
                code="duplicate_source_record_code",
                sheet=row.sheet,
                row=row.row,
                field_name="记录编码",
                value=code,
                message="来源记录编码在工作簿内重复",
                suggestion="每条来源记录使用唯一编码",
                planned_row=row,
            )
        seen_codes.add(code)
        row.existing = next((item for item in existing if item.record_code == code), None)
        current = None
        if row.existing:
            current = {
                "record_code": row.existing.record_code,
                "category": row.existing.category,
                "task_code": row.existing.task_code,
                "content": row.existing.content,
                "notes": row.existing.notes,
            }
        _set_status(row, current)
        result.append(row)
    return result


def _subject_current(subject: Subject) -> dict[str, Any]:
    return {
        "school_code": subject.school_code,
        "name": subject.name,
        "domain": subject.domain,
        "required_room_type": subject.required_room_type or "normal",
    }


def _teacher_current(teacher: Teacher) -> dict[str, Any]:
    return {
        "school_code": teacher.school_code,
        "name": teacher.name,
        "base_periods": teacher.base_periods,
        "admin_title": teacher.admin_title,
        "admin_reduction": teacher.admin_reduction,
        "arrangement_status": teacher.arrangement_status,
        "is_external": teacher.is_external,
        "is_active": teacher.is_active,
    }


def _room_current(room: Room) -> dict[str, Any]:
    return {
        "school_code": room.school_code,
        "name": room.name,
        "room_type": room.room_type,
        "capacity": room.capacity,
        "applicable_subjects": sorted(
            _object_reference(subject) for subject in room.subjects
        ),
    }


def _class_current(class_unit: ClassUnit) -> dict[str, Any]:
    return {
        "school_code": class_unit.school_code,
        "name": class_unit.name,
        "grade": class_unit.grade,
        "track": class_unit.track,
        "department": class_unit.department,
        "planned_weekly_periods": class_unit.planned_weekly_periods,
        "homeroom_teacher": (
            _object_reference(class_unit.homeroom_teacher)
            if class_unit.homeroom_teacher
            else None
        ),
        "period_table": (
            _object_reference(class_unit.period_table)
            if class_unit.period_table
            else None
        ),
    }


def _source_record_current(source: TeacherArrangementSourceRecord) -> dict[str, Any]:
    return {
        "record_code": source.record_code,
        "category": source.category,
        "task_code": source.task_code,
        "content": source.content,
        "notes": source.notes,
    }


def _target_current(entity: str, target: Any) -> dict[str, Any]:
    if entity == "subjects":
        return _subject_current(target)
    if entity == "teachers":
        return _teacher_current(target)
    if entity == "rooms":
        return _room_current(target)
    if entity == "period_tables":
        return _period_table_current(target)
    if entity == "classes":
        return _class_current(target)
    if entity == "assignments":
        return _assignment_current(target)
    if entity == "source_records":
        return _source_record_current(target)
    raise ValueError(f"不支持的教师安排实体：{entity}")


def _latest_import_records(
    db: Session, semester_id: int
) -> dict[tuple[str, str], TeacherArrangementImportRecord]:
    records = list(
        db.scalars(
            select(TeacherArrangementImportRecord)
            .where(TeacherArrangementImportRecord.semester_id == semester_id)
            .order_by(TeacherArrangementImportRecord.id.desc())
        )
    )
    latest: dict[tuple[str, str], TeacherArrangementImportRecord] = {}
    for record in records:
        latest.setdefault((record.entity_type, record.source_key), record)
    return latest


def _record_target(
    db: Session, record: TeacherArrangementImportRecord
) -> Any | None:
    model_by_entity = {
        "subjects": Subject,
        "teachers": Teacher,
        "rooms": Room,
        "period_tables": PeriodTable,
        "classes": ClassUnit,
        "assignments": CourseAssignment,
        "source_records": TeacherArrangementSourceRecord,
    }
    model = model_by_entity.get(record.entity_type)
    return db.get(model, record.target_id) if model is not None else None


def _decision_issue(
    issues: list[LocatedIssue], row: PlannedRow, selected: str, options: tuple[str, ...]
) -> None:
    _issue(
        issues,
        code="teacher_arrangement_decision_invalid",
        sheet=row.sheet,
        row=row.row,
        field_name="导入决策",
        value=selected,
        message="导入决策值无效",
        suggestion=f"请选择 {' / '.join(options)}",
        planned_row=row,
    )


def _reconcile_row(
    row: PlannedRow,
    record: TeacherArrangementImportRecord,
    decisions: dict[str, str],
    issues: list[LocatedIssue],
) -> None:
    if record.outcome != "applied" or row.existing is None:
        return
    if row.existing.id != record.target_id or row.current_values is None:
        return
    row.prior_record = record
    row.baseline_values = dict(record.applied_values)
    fields_to_apply: set[str] = set()
    conflict_fields: set[str] = set()
    changes: list[dict[str, Any]] = []
    for key, incoming in row.values.items():
        baseline = row.baseline_values.get(key)
        current = row.current_values.get(key)
        baseline_json = _jsonable(baseline)
        current_json = _jsonable(current)
        incoming_json = _jsonable(incoming)
        workbook_changed = incoming_json != baseline_json
        system_changed = current_json != baseline_json
        if not workbook_changed and not system_changed:
            continue
        if current_json == incoming_json:
            resolution = "same"
        elif workbook_changed and not system_changed:
            resolution = "workbook"
            fields_to_apply.add(key)
        elif system_changed and not workbook_changed:
            resolution = "system"
        else:
            resolution = "conflict"
            conflict_fields.add(key)
        changes.append(
            {
                "field": key,
                "before": current_json,
                "after": incoming_json if resolution == "workbook" else current_json,
                "baseline": baseline_json,
                "current": current_json,
                "incoming": incoming_json,
                "resolution": resolution,
            }
        )

    row.changes = changes
    if conflict_fields:
        row.decision_kind = "conflict"
        selected = decisions.get(row.decision_key)
        if selected is not None and selected not in {"incoming", "current"}:
            _decision_issue(issues, row, selected, ("incoming", "current"))
            selected = None
        row.decision_selected = selected
        if selected == "incoming":
            fields_to_apply.update(conflict_fields)
            for change in row.changes:
                if change["field"] in conflict_fields:
                    change["after"] = change["incoming"]
                    change["resolution"] = "incoming"
        elif selected == "current":
            for change in row.changes:
                if change["field"] in conflict_fields:
                    change["resolution"] = "current"
        else:
            row.fields_to_apply = fields_to_apply
            row.status = "conflict"
            return

    row.fields_to_apply = fields_to_apply
    row.status = "changed" if fields_to_apply else "unchanged"


_REFERENCE_LABELS = {
    "assignment_teachers": "教学任务",
    "course_assignments": "教学任务",
    "scheduling_unit_members": "排课单位",
    "timetable_entries": "课表",
    "teacher_time_rules": "教师时段规则",
    "teacher_subjects": "教师任教科目",
    "room_subjects": "教室适用科目",
    "class_units": "班主任关系",
    "leave_requests": "请假记录",
    "affected_periods": "调课与代课记录",
    "notifications": "通知记录",
    "periods": "作息节次",
}


def _removal_safety(db: Session, entity: str, target: Any | None) -> tuple[bool, str | None]:
    if target is None or entity == "source_records":
        return True, None
    owned_children = {
        "assignments": {"assignment_teachers", "block_rules"},
        "period_tables": {"periods"},
    }.get(entity, set())
    target_table = target.__table__
    for table in Base.metadata.tables.values():
        if table.name in owned_children:
            continue
        for foreign_key in table.foreign_keys:
            if foreign_key.column.table is not target_table:
                continue
            referenced = db.execute(
                select(foreign_key.parent)
                .where(foreign_key.parent == target.id)
                .limit(1)
            ).first()
            if referenced is not None:
                label = _REFERENCE_LABELS.get(table.name, table.name)
                return False, f"仍被{label}引用，只能保留"
    return True, None


def _add_disappeared_rows(
    db: Session,
    rows: dict[str, list[PlannedRow]],
    latest_records: dict[tuple[str, str], TeacherArrangementImportRecord],
    decisions: dict[str, str],
    issues: list[LocatedIssue],
) -> None:
    current_keys = {
        (entity, row.source_key)
        for entity, entity_rows in rows.items()
        for row in entity_rows
    }
    for (entity, source_key), record in latest_records.items():
        if entity not in rows or (entity, source_key) in current_keys:
            continue
        if record.outcome != "applied":
            continue
        target = _record_target(db, record)
        current_values = _target_current(entity, target) if target is not None else None
        removal_allowed, removal_reason = _removal_safety(db, entity, target)
        row = PlannedRow(
            entity=entity,
            sheet=record.sheet_name,
            row=record.row_number,
            source_key=source_key,
            identity=source_key.removeprefix("code:").removeprefix("name:"),
            values=dict(record.applied_values),
            raw_values=dict(record.raw_values),
            status="disappeared",
            current_values=current_values,
            baseline_values=dict(record.applied_values),
            decision_kind="disappeared",
            removal_allowed=removal_allowed,
            removal_reason=removal_reason,
            existing=target,
            prior_record=record,
        )
        if current_values is not None and _jsonable(current_values) != _jsonable(
            record.applied_values
        ):
            row.changes = [
                {
                    "field": key,
                    "before": _jsonable(record.applied_values.get(key)),
                    "after": _jsonable(current_values.get(key)),
                    "baseline": _jsonable(record.applied_values.get(key)),
                    "current": _jsonable(current_values.get(key)),
                    "incoming": None,
                    "resolution": "disappeared",
                }
                for key in record.applied_values
                if _jsonable(record.applied_values.get(key))
                != _jsonable(current_values.get(key))
            ]
        selected = decisions.get(row.decision_key)
        if selected is not None and selected not in {"keep", "remove"}:
            _decision_issue(issues, row, selected, ("keep", "remove"))
            selected = None
        row.decision_selected = selected
        if selected == "remove" and not removal_allowed:
            _issue(
                issues,
                code="teacher_arrangement_remove_unsafe",
                sheet=row.sheet,
                row=row.row,
                field_name="移除决策",
                value=row.identity,
                message="该来源数据当前不能安全移除",
                suggestion=removal_reason or "请保留该项并先解除关联",
                planned_row=row,
            )
        rows[entity].append(row)


def _reconcile_import_history(
    db: Session,
    semester_id: int,
    rows: dict[str, list[PlannedRow]],
    decisions: dict[str, str],
    issues: list[LocatedIssue],
) -> None:
    latest_records = _latest_import_records(db, semester_id)
    for entity, entity_rows in rows.items():
        for row in entity_rows:
            record = latest_records.get((entity, row.source_key))
            if record is not None:
                _reconcile_row(row, record, decisions, issues)
    _add_disappeared_rows(db, rows, latest_records, decisions, issues)
    allowed_keys = {
        row.decision_key
        for entity_rows in rows.values()
        for row in entity_rows
        if row.decision_kind is not None
    }
    for unknown_key in sorted(set(decisions) - allowed_keys):
        _issue(
            issues,
            code="teacher_arrangement_decision_unknown",
            sheet="导入决策",
            row=0,
            field_name="决策键",
            value=unknown_key,
            message="决策项不属于当前预览",
            suggestion="请重新预览并只提交当前显示的决策",
        )


def _database_snapshot(db: Session, semester_id: int) -> dict[str, Any]:
    subjects = list(db.scalars(select(Subject).where(Subject.semester_id == semester_id)))
    teachers = list(db.scalars(select(Teacher).where(Teacher.semester_id == semester_id)))
    rooms = list(db.scalars(select(Room).where(Room.semester_id == semester_id)))
    period_tables = list(
        db.scalars(select(PeriodTable).where(PeriodTable.semester_id == semester_id))
    )
    classes = list(db.scalars(select(ClassUnit).where(ClassUnit.semester_id == semester_id)))
    assignments = list(
        db.scalars(select(CourseAssignment).where(CourseAssignment.semester_id == semester_id))
    )
    source_records = list(
        db.scalars(
            select(TeacherArrangementSourceRecord).where(
                TeacherArrangementSourceRecord.semester_id == semester_id
            )
        )
    )
    import_history = _latest_import_records(db, semester_id)
    return {
        "subjects": [
            [
                item.id,
                item.school_code,
                item.name,
                item.domain,
                item.required_room_type,
            ]
            for item in sorted(subjects, key=lambda item: item.id)
        ],
        "teachers": [
            [
                item.id,
                item.school_code,
                item.name,
                item.base_periods,
                item.admin_title,
                item.admin_reduction,
                item.arrangement_status,
                item.is_external,
                item.is_active,
            ]
            for item in sorted(teachers, key=lambda item: item.id)
        ],
        "rooms": [
            [item.id, *_room_current(item).values()]
            for item in sorted(rooms, key=lambda item: item.id)
        ],
        "period_tables": [
            [item.id, *_period_table_current(item).values()]
            for item in sorted(period_tables, key=lambda item: item.id)
        ],
        "classes": [
            [
                item.id,
                item.school_code,
                item.name,
                item.grade,
                item.track,
                item.department,
                item.planned_weekly_periods,
                item.homeroom_teacher_id,
                item.period_table_id,
            ]
            for item in sorted(classes, key=lambda item: item.id)
        ],
        "assignments": [
            [item.id, *_assignment_current(item).values()]
            for item in sorted(assignments, key=lambda item: item.id)
        ],
        "source_records": [
            [
                item.id,
                item.record_code,
                item.category,
                item.task_code,
                item.content,
                item.notes,
            ]
            for item in sorted(source_records, key=lambda item: item.id)
        ],
        "teacher_arrangement_history": [
            [
                record.id,
                record.batch_id,
                record.entity_type,
                record.source_key,
                record.target_id,
                record.outcome,
                record.applied_values,
            ]
            for record in sorted(import_history.values(), key=lambda item: item.id)
        ],
    }


def _fingerprint(
    *,
    workbook_sha256: str,
    semester_id: int,
    mode: TemplateMode,
    database_snapshot: dict[str, Any],
    decisions: dict[str, str],
) -> str:
    payload = {
        "plan_version": 4,
        "template_version": TEMPLATE_VERSION,
        "workbook_sha256": workbook_sha256,
        "semester_id": semester_id,
        "mode": mode,
        "database": database_snapshot,
        "decisions": decisions,
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode()).hexdigest()


def build_plan(
    db: Session,
    semester: Semester,
    content: bytes,
    mode: TemplateMode,
    decisions: dict[str, str] | None = None,
) -> ImportPlan:
    normalized_decisions = dict(sorted((decisions or {}).items()))
    workbook = _load_workbook(content, semester, mode)
    issues: list[LocatedIssue] = []
    raw_by_entity: dict[str, list[RawRow]] = {}
    for definition in visible_definitions(mode):
        raw_by_entity[definition.key] = _read_sheet_rows(
            workbook, definition, mode, issues
        )

    semester_rows = raw_by_entity.get("semester", [])
    if len(semester_rows) != 1:
        _issue(
            issues,
            code="semester_guard_row_invalid",
            sheet="学期",
            row=4,
            field_name="学期",
            value=len(semester_rows),
            message="学期表必须且只能保留一行目标学期",
            suggestion="请重新下载模板，不要新增或删除学期行",
        )
    elif (
        _integer(semester_rows[0], "academic_year", issues, minimum=1900, maximum=2100)
        != semester.academic_year
        or _integer(semester_rows[0], "term", issues, minimum=1, maximum=2)
        != semester.term
    ):
        _issue(
            issues,
            code="semester_guard_mismatch",
            sheet="学期",
            row=semester_rows[0].row,
            field_name="学年起始年",
            value={
                "academic_year": semester_rows[0].values.get("academic_year"),
                "term": semester_rows[0].values.get("term"),
            },
            message="工作簿学期与当前目标学期不一致",
            suggestion=(
                f"请保持 {semester.academic_year}-{semester.academic_year + 1} 学年"
                f"第 {semester.term} 学期"
            ),
        )

    rows: dict[str, list[PlannedRow]] = {key: [] for key in _entity_keys(mode)}
    rows["subjects"] = _plan_subjects(
        db, semester.id, raw_by_entity.get("subjects", []), issues
    )
    rows["teachers"] = _plan_teachers(
        db, semester.id, raw_by_entity.get("teachers", []), issues
    )
    if mode == "scheduling_ready":
        rows["rooms"] = _plan_rooms(
            db,
            semester.id,
            raw_by_entity.get("rooms", []),
            rows["subjects"],
            issues,
        )
        rows["period_tables"] = _plan_period_tables(
            db,
            semester.id,
            raw_by_entity.get("period_tables", []),
            issues,
        )
    rows["classes"] = _plan_classes(
        db,
        semester.id,
        raw_by_entity.get("classes", []),
        rows["teachers"],
        rows.get("period_tables", []),
        issues,
        mode,
    )
    rows["assignments"] = _plan_assignments(
        db,
        semester.id,
        raw_by_entity.get("assignments", []),
        rows["subjects"],
        rows["teachers"],
        rows["classes"],
        issues,
        mode,
    )
    rows["source_records"] = _plan_source_records(
        db, semester.id, raw_by_entity.get("source_records", []), issues
    )
    _reconcile_import_history(db, semester.id, rows, normalized_decisions, issues)
    if mode == "scheduling_ready":
        _validate_ready_plan(db, semester.id, rows, issues)
    workbook_sha256 = hashlib.sha256(content).hexdigest()
    return ImportPlan(
        semester_id=semester.id,
        mode=mode,
        workbook_sha256=workbook_sha256,
        fingerprint=_fingerprint(
            workbook_sha256=workbook_sha256,
            semester_id=semester.id,
            mode=mode,
            database_snapshot=_database_snapshot(db, semester.id),
            decisions=normalized_decisions,
        ),
        decisions=normalized_decisions,
        rows=rows,
        issues=issues,
    )


def _applied_target(target: PlannedRow | Any | None) -> Any:
    if isinstance(target, PlannedRow):
        return target.applied
    return target


def _apply_subjects(db: Session, semester_id: int, rows: list[PlannedRow]) -> None:
    for row in rows:
        if row.status == "disappeared":
            continue
        subject = row.existing or Subject(semester_id=semester_id)
        if row.existing is None:
            db.add(subject)
        for key in ("school_code", "name", "domain", "required_room_type"):
            if key in row.fields_to_apply:
                setattr(subject, key, row.values[key])
        row.applied = subject
    db.flush()


def _apply_teachers(db: Session, semester_id: int, rows: list[PlannedRow]) -> None:
    for row in rows:
        if row.status == "disappeared":
            continue
        teacher = row.existing or Teacher(semester_id=semester_id)
        if row.existing is None:
            db.add(teacher)
        for key in (
            "school_code",
            "name",
            "base_periods",
            "admin_title",
            "admin_reduction",
            "arrangement_status",
            "is_external",
            "is_active",
        ):
            if key in row.fields_to_apply:
                setattr(teacher, key, row.values[key])
        row.applied = teacher
    db.flush()


def _apply_rooms(db: Session, semester_id: int, rows: list[PlannedRow]) -> None:
    for row in rows:
        if row.status == "disappeared":
            continue
        room = row.existing or Room(semester_id=semester_id)
        if row.existing is None:
            db.add(room)
        for key in ("school_code", "name", "room_type", "capacity"):
            if key in row.fields_to_apply:
                setattr(room, key, row.values[key])
        if "applicable_subjects" in row.fields_to_apply:
            room.subjects = [
                _applied_target(target)
                for target in row.references.get("applicable_subjects", [])
            ]
        row.applied = room
    db.flush()


def _apply_period_tables(
    db: Session,
    semester_id: int,
    rows: list[PlannedRow],
) -> None:
    for row in rows:
        if row.status == "disappeared":
            continue
        table = row.existing or PeriodTable(semester_id=semester_id)
        if row.existing is None:
            db.add(table)
        for key in ("school_code", "name", "num_weekdays", "is_default"):
            if key in row.fields_to_apply:
                setattr(table, key, row.values[key])
        if "periods" in row.fields_to_apply:
            if row.existing is not None:
                table.periods.clear()
                db.flush()
            table.periods = [
                Period(
                    weekday=period["weekday"],
                    period_no=period["period_number"],
                    name=_period_name(
                        period["period_type"],
                        period["period_number"],
                    ),
                    type=period["period_type"],
                    start_time=_time_from_iso(period["start_time"]),
                    end_time=_time_from_iso(period["end_time"]),
                )
                for period in row.values["periods"]
            ]
        row.applied = table
    db.flush()


def _apply_classes(db: Session, semester_id: int, rows: list[PlannedRow]) -> None:
    for row in rows:
        if row.status == "disappeared":
            continue
        class_unit = row.existing or ClassUnit(semester_id=semester_id)
        if row.existing is None:
            db.add(class_unit)
        for key in (
            "school_code",
            "name",
            "grade",
            "track",
            "department",
            "planned_weekly_periods",
        ):
            if key in row.fields_to_apply:
                setattr(class_unit, key, row.values[key])
        if "homeroom_teacher" in row.fields_to_apply:
            homeroom_teacher = _applied_target(row.references.get("homeroom_teacher"))
            class_unit.homeroom_teacher = homeroom_teacher
        if "period_table" in row.fields_to_apply:
            class_unit.period_table = _applied_target(row.references.get("period_table"))
        row.applied = class_unit
    db.flush()


def _apply_assignments(db: Session, semester_id: int, rows: list[PlannedRow]) -> None:
    for row in rows:
        if row.status == "disappeared":
            continue
        assignment = row.existing or CourseAssignment(semester_id=semester_id)
        if "task_code" in row.fields_to_apply:
            assignment.task_code = row.values["task_code"]
        if "component" in row.fields_to_apply:
            assignment.component = row.values["component"]
        if "class_ref" in row.fields_to_apply:
            class_unit = _applied_target(row.references["class"])
            unit = get_or_create_single_unit(db, class_unit)
            unit.name = class_unit.name
            assignment.scheduling_unit = unit
        if "subject_ref" in row.fields_to_apply:
            subject = _applied_target(row.references["subject"])
            assignment.subject = subject
        if "weekly_periods" in row.fields_to_apply:
            assignment.periods_per_week = row.values["weekly_periods"]
        if {"lead_teacher", "co_teachers"} & row.fields_to_apply:
            lead = _applied_target(row.references.get("lead_teacher"))
            desired_links: list[tuple[Teacher, bool]] = []
            if lead is not None:
                desired_links.append((lead, True))
            for teacher_target in row.references.get("co_teachers", []):
                teacher = _applied_target(teacher_target)
                desired_links.append((teacher, False))
            desired_by_teacher_id = {
                teacher.id: (teacher, is_lead) for teacher, is_lead in desired_links
            }
            existing_by_teacher_id = {
                link.teacher_id: link
                for link in assignment.teachers
                if link.teacher_id is not None
            }
            for link in list(assignment.teachers):
                desired = desired_by_teacher_id.get(link.teacher_id)
                if desired is None:
                    assignment.teachers.remove(link)
                else:
                    link.is_lead = desired[1]
            for teacher_id, (teacher, is_lead) in desired_by_teacher_id.items():
                if teacher_id not in existing_by_teacher_id:
                    assignment.teachers.append(
                        AssignmentTeacher(teacher=teacher, is_lead=is_lead)
                    )
        if row.existing is None:
            db.add(assignment)
        assignment.required_room_type = (
            None
            if assignment.subject.required_room_type == RoomType.normal.value
            else assignment.subject.required_room_type
        )
        row.applied = assignment
    db.flush()


def _apply_source_records(db: Session, semester_id: int, rows: list[PlannedRow]) -> None:
    for row in rows:
        if row.status == "disappeared":
            continue
        source = row.existing or TeacherArrangementSourceRecord(semester_id=semester_id)
        if row.existing is None:
            db.add(source)
        for key in ("record_code", "category", "task_code", "content", "notes"):
            if key in row.fields_to_apply:
                setattr(source, key, row.values[key])
        row.applied = source
    db.flush()


def _apply_disappeared(db: Session, plan: ImportPlan) -> None:
    removal_order = (
        "assignments",
        "source_records",
        "classes",
        "period_tables",
        "rooms",
        "teachers",
        "subjects",
    )
    for entity in removal_order:
        for row in plan.rows.get(entity, []):
            if row.status != "disappeared":
                continue
            if row.decision_selected == "keep":
                row.applied = row.existing
                continue
            if row.decision_selected != "remove":
                raise ValueError("消失来源尚未完成保留或移除决策")
            if not row.removal_allowed:
                raise ValueError("消失来源当前不能安全移除")
            if row.existing is not None:
                db.delete(row.existing)
        db.flush()


def _result_counts(
    plan: ImportPlan,
) -> tuple[
    dict[str, int],
    dict[str, int],
    dict[str, int],
    dict[str, int],
    dict[str, int],
]:
    created = {key: 0 for key in plan.rows}
    updated = {key: 0 for key in plan.rows}
    unchanged = {key: 0 for key in plan.rows}
    removed = {key: 0 for key in plan.rows}
    kept = {key: 0 for key in plan.rows}
    for entity, rows in plan.rows.items():
        for row in rows:
            if row.status == "new":
                created[entity] += 1
            elif row.status == "changed":
                updated[entity] += 1
            elif row.status == "disappeared":
                if row.decision_selected == "remove":
                    removed[entity] += 1
                elif row.decision_selected == "keep":
                    kept[entity] += 1
            else:
                unchanged[entity] += 1
    return created, updated, unchanged, removed, kept


def apply_plan(db: Session, plan: ImportPlan, *, filename: str) -> dict[str, Any]:
    if not plan.can_commit:
        raise ValueError("导入计划仍有阻断项")
    _apply_disappeared(db, plan)
    _apply_subjects(db, plan.semester_id, plan.rows["subjects"])
    _apply_teachers(db, plan.semester_id, plan.rows["teachers"])
    if "rooms" in plan.rows:
        _apply_rooms(db, plan.semester_id, plan.rows["rooms"])
    if "period_tables" in plan.rows:
        _apply_period_tables(db, plan.semester_id, plan.rows["period_tables"])
    _apply_classes(db, plan.semester_id, plan.rows["classes"])
    _apply_assignments(db, plan.semester_id, plan.rows["assignments"])
    _apply_source_records(db, plan.semester_id, plan.rows["source_records"])

    semester = db.get(Semester, plan.semester_id)
    if semester is not None:
        semester.readiness = SemesterReadiness.draft.value

    created, updated, unchanged, removed, kept = _result_counts(plan)
    batch = TeacherArrangementImportBatch(
        semester_id=plan.semester_id,
        fingerprint=plan.fingerprint,
        workbook_sha256=plan.workbook_sha256,
        template_version=TEMPLATE_VERSION,
        mode=plan.mode,
        filename=filename[:255],
        decisions=plan.decisions,
    )
    db.add(batch)
    db.flush()
    summary = {
        "created": created,
        "updated": updated,
        "unchanged": unchanged,
        "removed": removed,
        "kept": kept,
        "fingerprint": plan.fingerprint,
    }
    batch.summary = summary
    for entity, rows in plan.rows.items():
        for row in rows:
            outcome = "applied"
            target_id = row.applied.id if row.applied is not None else None
            applied_values = row.effective_values()
            if row.status == "disappeared":
                outcome = "removed" if row.decision_selected == "remove" else "kept"
                target_id = row.prior_record.target_id if row.prior_record else target_id
                applied_values = (
                    {}
                    if outcome == "removed"
                    else _jsonable(row.current_values or row.baseline_values or {})
                )
            if target_id is None:
                raise ValueError("导入记录缺少目标标识")
            db.add(
                TeacherArrangementImportRecord(
                    batch_id=batch.id,
                    semester_id=plan.semester_id,
                    entity_type=entity,
                    source_key=row.source_key,
                    sheet_name=row.sheet,
                    row_number=row.row,
                    target_id=target_id,
                    outcome=outcome,
                    applied_values=applied_values,
                    raw_values=_jsonable(row.raw_values),
                )
            )
    db.flush()
    return {
        "batch_id": batch.id,
        **summary,
        "idempotent": False,
    }


def find_committed_batch(
    db: Session,
    semester_id: int,
    fingerprint: str,
    workbook_sha256: str,
) -> TeacherArrangementImportBatch | None:
    return db.scalar(
        select(TeacherArrangementImportBatch).where(
            TeacherArrangementImportBatch.semester_id == semester_id,
            TeacherArrangementImportBatch.fingerprint == fingerprint,
            TeacherArrangementImportBatch.workbook_sha256 == workbook_sha256,
        )
    )


def idempotent_result(batch: TeacherArrangementImportBatch) -> dict[str, Any]:
    return {"batch_id": batch.id, **batch.summary, "idempotent": True}
