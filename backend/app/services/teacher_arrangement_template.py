"""版本化教师安排模板生成。

中文表头服务于填表人员；隐藏 schema 中的稳定键服务于后续导入器。模板格式
一经发布只能通过新增版本演进，不能依赖表头文字作为持久身份。
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.models.semester import Semester

TEMPLATE_TYPE = "teacher_arrangement"
TEMPLATE_VERSION = "1.0"
MAX_DATA_ROW = 5000
TemplateMode = Literal["standard", "scheduling_ready"]

ROOM_TYPES = ("普通教室", "专用教室", "实训场地", "户外场地")
TEACHER_STATUSES = ("在岗", "外出", "产假", "停用")
YES_NO = ("是", "否")
TRACKS = ("小学", "初中", "普通高中", "综合高中", "中职", "职业高中")
WEEKDAYS = ("星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日")
PERIOD_TYPES = ("常规课时", "晨会", "午休", "课间", "活动")


@dataclass(frozen=True, slots=True)
class FieldDefinition:
    key: str
    header: str
    instruction: str
    example: object
    aliases: tuple[str, ...] = ()
    required: bool = False
    ready_only: bool = False
    choices: tuple[str, ...] = ()
    integer_range: tuple[int, int] | None = None
    time_value: bool = False


@dataclass(frozen=True, slots=True)
class SheetDefinition:
    key: str
    label: str
    fields: tuple[FieldDefinition, ...]
    ready_only: bool = False


SHEET_DEFINITIONS = (
    SheetDefinition(
        "semester",
        "学期",
        (
            FieldDefinition(
                "academic_year",
                "学年起始年",
                "必填；填写四位数起始年份，例如 2026 表示 2026-2027 学年",
                2025,
                aliases=("学年", "学年开始年份"),
                required=True,
                integer_range=(1900, 2100),
            ),
            FieldDefinition(
                "term",
                "学期",
                "必填；填写 1 或 2",
                1,
                aliases=("学期序号",),
                required=True,
                choices=("1", "2"),
            ),
        ),
    ),
    SheetDefinition(
        "subjects",
        "科目",
        (
            FieldDefinition(
                "school_code",
                "学校科目编码",
                "选填；学校已有稳定编码时填写，编码一旦使用将优先于名称识别",
                "SUB-MATH",
                aliases=("科目编码",),
            ),
            FieldDefinition(
                "name",
                "科目名称",
                "必填；无学校编码时须在目标学期内唯一",
                "数学",
                aliases=("名称", "课程名称"),
                required=True,
            ),
            FieldDefinition(
                "domain",
                "领域/类别",
                "选填；用于基础数据分类",
                "数学",
                aliases=("类别", "学科领域"),
            ),
            FieldDefinition(
                "required_room_type",
                "所需教室/场地类型",
                "选填；不填按普通教室处理",
                "普通教室",
                aliases=("所需教室类型", "教室类型"),
                choices=ROOM_TYPES,
            ),
        ),
    ),
    SheetDefinition(
        "teachers",
        "教师",
        (
            FieldDefinition(
                "school_code",
                "学校教师编码",
                "选填；学校已有稳定工号时填写，编码一旦使用将优先于姓名识别",
                "T-001",
                aliases=("教师编码", "工号"),
            ),
            FieldDefinition(
                "name",
                "教师姓名",
                "必填；无学校编码时须在目标学期内唯一",
                "王老师",
                aliases=("姓名",),
                required=True,
            ),
            FieldDefinition(
                "base_periods",
                "基础周课时",
                "选填；非负整数，不填按 0 导入并给出警告",
                18,
                aliases=("基本课时", "基础课时"),
                integer_range=(0, 100),
            ),
            FieldDefinition(
                "admin_title",
                "行政职务",
                "选填；只记录结构化职务，不从备注推断",
                "年级负责人",
                aliases=("职务",),
            ),
            FieldDefinition(
                "admin_reduction",
                "行政减课时",
                "选填；非负整数，不填按 0 导入并给出警告",
                2,
                aliases=("行政减课", "减课时"),
                integer_range=(0, 100),
            ),
            FieldDefinition(
                "status",
                "教师状态",
                "必填；外出、产假等必须使用结构化状态",
                "在岗",
                aliases=("状态",),
                required=True,
                choices=TEACHER_STATUSES,
            ),
            FieldDefinition(
                "is_external",
                "外聘",
                "选填；填写是或否，不填按否处理",
                "否",
                aliases=("外聘教师",),
                choices=YES_NO,
            ),
        ),
    ),
    SheetDefinition(
        "classes",
        "班级",
        (
            FieldDefinition(
                "school_code",
                "学校班级编码",
                "选填；学校已有稳定编码时填写，编码一旦使用将优先于名称识别",
                "C-701",
                aliases=("班级编码",),
            ),
            FieldDefinition(
                "name",
                "班级名称",
                "必填；无学校编码时须在目标学期内唯一",
                "七年级1班",
                aliases=("班名",),
                required=True,
            ),
            FieldDefinition(
                "grade",
                "年级",
                "必填；填写 1-12 的整数",
                7,
                required=True,
                integer_range=(1, 12),
            ),
            FieldDefinition(
                "track",
                "学制",
                "必填；从下拉选项中选择",
                "初中",
                aliases=("培养类型",),
                required=True,
                choices=TRACKS,
            ),
            FieldDefinition(
                "specialization",
                "专业/班级类别",
                "选填；职业或综合高中可填写专业方向",
                "普通班",
                aliases=("专业", "班级类别"),
            ),
            FieldDefinition(
                "homeroom_teacher",
                "班主任",
                "选填；填写学校教师编码或唯一教师姓名",
                "王老师",
                aliases=("班主任教师",),
            ),
            FieldDefinition(
                "planned_weekly_periods",
                "班级计划周课时",
                "必填；该班所有可排教学任务的周课时目标总数",
                35,
                aliases=("计划周课时",),
                required=True,
                integer_range=(0, 200),
            ),
            FieldDefinition(
                "period_table_code",
                "作息表编码",
                "排课准备模式必填；引用作息时间表中的编码",
                "PT-JUNIOR",
                aliases=("作息时间表编码",),
                required=True,
                ready_only=True,
            ),
        ),
    ),
    SheetDefinition(
        "assignments",
        "教学任务",
        (
            FieldDefinition(
                "task_code",
                "任务编码",
                "必填且在目标学期内唯一；重导入依靠该编码识别任务",
                "TASK-701-MATH",
                aliases=("教学任务编码",),
                required=True,
            ),
            FieldDefinition(
                "class_ref",
                "班级",
                "必填；填写学校班级编码或唯一班级名称",
                "C-701",
                aliases=("班级编码/名称", "班名"),
                required=True,
            ),
            FieldDefinition(
                "subject_ref",
                "科目",
                "必填；填写学校科目编码或唯一科目名称",
                "SUB-MATH",
                aliases=("科目编码/名称", "课程"),
                required=True,
            ),
            FieldDefinition(
                "component",
                "组成",
                "必填；一行只表示一个可独立排课的组成，不得填写 4+1",
                "基础课",
                aliases=("任务组成", "课程组成"),
                required=True,
            ),
            FieldDefinition(
                "weekly_periods",
                "周课时",
                "必填；填写正整数，复合课时须拆成多行",
                5,
                aliases=("每周课时",),
                required=True,
                integer_range=(1, 40),
            ),
            FieldDefinition(
                "lead_teacher",
                "主讲教师",
                "标准模式可空并给出警告；排课准备模式必填",
                "T-001",
                aliases=("任课教师", "教师"),
            ),
            FieldDefinition(
                "co_teachers",
                "协同教师",
                "选填；多名以顿号分隔，表示每节课共同到场",
                "T-002、T-003",
                aliases=("共同教师", "协作教师"),
            ),
        ),
    ),
    SheetDefinition(
        "source_records",
        "来源记录",
        (
            FieldDefinition(
                "record_code",
                "记录编码",
                "必填且唯一；用于追溯原教师安排文档中的非排课信息",
                "SRC-001",
                aliases=("来源编码",),
                required=True,
            ),
            FieldDefinition(
                "category",
                "记录类型",
                "必填；例如费用、人数、缺编、周期或备注",
                "备注",
                aliases=("类型",),
                required=True,
            ),
            FieldDefinition(
                "task_code",
                "关联任务编码",
                "选填；需要时关联教学任务，不会据此创建任务",
                "TASK-701-MATH",
                aliases=("任务编码",),
            ),
            FieldDefinition(
                "content",
                "原始内容",
                "必填；忠实记录来源文档内容，不参与自动排课",
                "数学教师缺编 1 人",
                aliases=("原文", "内容"),
                required=True,
            ),
            FieldDefinition(
                "notes",
                "备注",
                "选填；补充人工整理说明",
                "来源于教师安排表备注栏",
                aliases=("整理说明",),
            ),
        ),
    ),
    SheetDefinition(
        "rooms",
        "教室及户外场地",
        (
            FieldDefinition(
                "school_code",
                "学校教室/场地编码",
                "选填；学校已有稳定编码时填写",
                "ROOM-LAB-A",
                aliases=("教室编码", "户外场地编码"),
            ),
            FieldDefinition(
                "name",
                "教室/场地名称",
                "必填；无学校编码时须在目标学期内唯一",
                "物理实验室A",
                aliases=("名称", "教室名称", "户外场地名称"),
                required=True,
            ),
            FieldDefinition(
                "room_type",
                "教室/场地类型",
                "必填；从下拉选项中选择",
                "专用教室",
                aliases=("类型", "教室类型", "户外场地类型"),
                required=True,
                choices=ROOM_TYPES,
            ),
            FieldDefinition(
                "capacity",
                "容量",
                "选填；非负整数",
                48,
                aliases=("容纳人数",),
                integer_range=(0, 5000),
            ),
            FieldDefinition(
                "applicable_subjects",
                "适用科目",
                "选填；多科目以顿号分隔，填写学校科目编码或唯一名称",
                "SUB-PHYSICS、SUB-CHEMISTRY",
                aliases=("适用课程",),
            ),
        ),
        ready_only=True,
    ),
    SheetDefinition(
        "period_tables",
        "作息时间表",
        (
            FieldDefinition(
                "table_code",
                "作息表编码",
                "必填；同一编码的多行共同组成一套作息表",
                "PT-JUNIOR",
                aliases=("作息时间表编码",),
                required=True,
            ),
            FieldDefinition(
                "table_name",
                "作息表名称",
                "必填；同一编码各行名称必须一致",
                "初中部作息",
                aliases=("作息时间表名称",),
                required=True,
            ),
            FieldDefinition(
                "weekday",
                "星期",
                "必填；从下拉选项中选择",
                "星期一",
                aliases=("周几",),
                required=True,
                choices=WEEKDAYS,
            ),
            FieldDefinition(
                "period_number",
                "节次",
                "必填；填写正整数",
                1,
                aliases=("第几节",),
                required=True,
                integer_range=(1, 20),
            ),
            FieldDefinition(
                "period_type",
                "课时类型",
                "必填；只有常规课时计入可排容量",
                "常规课时",
                aliases=("类型",),
                required=True,
                choices=PERIOD_TYPES,
            ),
            FieldDefinition(
                "start_time",
                "开始时间",
                "必填；使用 Excel 时间格式，例如 08:00",
                "08:00",
                aliases=("上课时间",),
                required=True,
                time_value=True,
            ),
            FieldDefinition(
                "end_time",
                "结束时间",
                "必填；使用 Excel 时间格式，例如 08:45",
                "08:45",
                aliases=("下课时间",),
                required=True,
                time_value=True,
            ),
        ),
        ready_only=True,
    ),
)

_HEADER_FILL = PatternFill("solid", fgColor="284B63")
_INSTRUCTION_FILL = PatternFill("solid", fgColor="E8F0F4")
_EXAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")
_TARGET_FILL = PatternFill("solid", fgColor="E2F0D9")
_THIN_BORDER = Border(bottom=Side(style="thin", color="B8C2CC"))


def visible_definitions(mode: TemplateMode) -> tuple[SheetDefinition, ...]:
    return tuple(
        definition
        for definition in SHEET_DEFINITIONS
        if mode == "scheduling_ready" or not definition.ready_only
    )


def fields_for_mode(
    definition: SheetDefinition, mode: TemplateMode
) -> tuple[FieldDefinition, ...]:
    return tuple(
        field
        for field in definition.fields
        if mode == "scheduling_ready" or not field.ready_only
    )


def _add_list_validation(sheet, column: str, values: tuple[str, ...], allow_blank: bool) -> None:
    validation = DataValidation(
        type="list",
        formula1=f'"{",".join(values)}"',
        allow_blank=allow_blank,
        error="请从下拉列表中选择有效值",
        errorTitle="无效选项",
    )
    validation.errorStyle = "stop"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{column}4:{column}{MAX_DATA_ROW}")


def _add_integer_validation(
    sheet, column: str, bounds: tuple[int, int], allow_blank: bool
) -> None:
    validation = DataValidation(
        type="whole",
        operator="between",
        formula1=str(bounds[0]),
        formula2=str(bounds[1]),
        allow_blank=allow_blank,
        error=f"请填写 {bounds[0]} 至 {bounds[1]} 之间的整数",
        errorTitle="无效整数",
    )
    validation.errorStyle = "stop"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{column}4:{column}{MAX_DATA_ROW}")


def _add_time_validation(sheet, column: str, allow_blank: bool) -> None:
    validation = DataValidation(
        type="time",
        operator="between",
        formula1="TIME(0,0,0)",
        formula2="TIME(23,59,59)",
        allow_blank=allow_blank,
        error="请填写有效的 Excel 时间，例如 08:00",
        errorTitle="无效时间",
    )
    validation.errorStyle = "stop"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{column}4:{column}{MAX_DATA_ROW}")


def _add_required_text_validation(sheet, column: str) -> None:
    validation = DataValidation(
        type="custom",
        formula1=f'LEN(TRIM({column}4&""))>0',
        allow_blank=False,
        error="该字段为必填项",
        errorTitle="缺少必填项",
    )
    validation.errorStyle = "stop"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{column}4:{column}{MAX_DATA_ROW}")


def _style_data_sheet(sheet, fields: tuple[FieldDefinition, ...]) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER
    for cell in sheet[2]:
        cell.fill = _INSTRUCTION_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[3]:
        cell.fill = _EXAMPLE_FILL
        cell.font = Font(color="7F6000", italic=True)
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 58
    sheet.row_dimensions[3].height = 24
    sheet.freeze_panes = "A4"
    last_column = get_column_letter(len(fields))
    sheet.auto_filter.ref = f"A1:{last_column}{MAX_DATA_ROW}"
    sheet.sheet_view.showGridLines = False
    for index, field in enumerate(fields, start=1):
        column = get_column_letter(index)
        longest = max(len(field.header), len(field.instruction), len(str(field.example)))
        sheet.column_dimensions[column].width = min(max(longest + 2, 14), 38)
        if field.choices:
            _add_list_validation(sheet, column, field.choices, not field.required)
        elif field.integer_range:
            _add_integer_validation(sheet, column, field.integer_range, not field.required)
        elif field.time_value:
            _add_time_validation(sheet, column, not field.required)
        elif field.required:
            _add_required_text_validation(sheet, column)


def _create_instructions(workbook: Workbook, semester: Semester, mode: TemplateMode) -> None:
    sheet = workbook.create_sheet("说明")
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "教师安排标准化导入模板"
    sheet["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = _HEADER_FILL
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 36
    rows = (
        ("模板版本", f"v{TEMPLATE_VERSION}"),
        ("目标学期", semester.label),
        ("导入模式", "自动排课准备" if mode == "scheduling_ready" else "教师安排标准"),
        (
            "填写规则",
            "各数据表第 1 行为表头、第 2 行为说明、第 3 行为脱敏示例；正式数据从第 4 行开始。",
        ),
        (
            "身份识别",
            "任务编码必填；学校编码选填。已有学校编码时编码优先，否则按目标学期内唯一名称识别。",
        ),
        (
            "课时拆分",
            "一行只能表示一个可独立排课的任务；4+1、3+2 等必须拆成多行并使用不同任务编码。",
        ),
        ("协同教师", "同一行协同教师表示每节课共同到场；拆分授课请建立多行任务。"),
        ("非排课信息", "费用、人数、缺编、周期和自由文本放入来源记录，不要创建虚假教学任务。"),
        ("导入顺序", "下载模板 -> 填写并上传 -> 修正阻断项 -> 确认导入；导入成功不等于排课就绪。"),
    )
    for row in rows:
        sheet.append(row)
    for cell in sheet["A"]:
        cell.font = Font(bold=True, color="284B63")
    for row in sheet.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _THIN_BORDER
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 92
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"


def _create_data_sheet(
    workbook: Workbook,
    definition: SheetDefinition,
    mode: TemplateMode,
    semester: Semester,
) -> None:
    fields = fields_for_mode(definition, mode)
    sheet = workbook.create_sheet(definition.label)
    sheet.append([field.header for field in fields])
    sheet.append([field.instruction for field in fields])
    sheet.append([field.example for field in fields])
    if definition.key == "semester":
        sheet.append([semester.academic_year, semester.term])
        for cell in sheet[4]:
            cell.fill = _TARGET_FILL
            cell.font = Font(bold=True, color="375623")
    _style_data_sheet(sheet, fields)


def _create_schema(workbook: Workbook, semester: Semester, mode: TemplateMode) -> None:
    sheet = workbook.create_sheet("_schema")
    metadata = (
        ("template_type", TEMPLATE_TYPE),
        ("template_version", TEMPLATE_VERSION),
        ("mode", mode),
        ("semester_id", semester.id),
        ("academic_year", semester.academic_year),
        ("term", semester.term),
        ("header_rows", 3),
    )
    for row in metadata:
        sheet.append(row)
    sheet.append([])
    sheet.append(
        ["sheet_key", "sheet_name", "field_key", "header", "aliases", "required"]
    )
    for definition in visible_definitions(mode):
        for field in fields_for_mode(definition, mode):
            sheet.append(
                [
                    definition.key,
                    definition.label,
                    field.key,
                    field.header,
                    json.dumps(field.aliases, ensure_ascii=False),
                    field.required,
                ]
            )
    sheet.sheet_state = "veryHidden"


def build_template(semester: Semester, mode: TemplateMode) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "教师安排标准化导入模板"
    workbook.properties.subject = f"{semester.label} / {mode} / v{TEMPLATE_VERSION}"
    _create_instructions(workbook, semester, mode)
    for definition in visible_definitions(mode):
        _create_data_sheet(workbook, definition, mode, semester)
    _create_schema(workbook, semester, mode)
    workbook.active = 0
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def download_filename(semester: Semester, mode: TemplateMode) -> str:
    return (
        f"teacher_arrangement_{semester.academic_year}-{semester.academic_year + 1}"
        f"_term-{semester.term}_{mode}_v{TEMPLATE_VERSION}.xlsx"
    )
