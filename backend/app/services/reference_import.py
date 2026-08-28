"""Word/Excel 参考文件适配器。

这不是通用 Excel 导入器：两份文件的版式是学校内部工作表，适配器先将其
归一化为可审阅的导入计划，再由 API 在一个事务中创建基础数据、教学任务和
固定课位。原文件不会被改写，已有人工数据也不会被静默覆盖。
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.assignment import AssignmentTeacher, CourseAssignment
from app.models.basedata import (
    ClassTrack,
    ClassUnit,
    Room,
    RoomType,
    Subject,
    Teacher,
    TeacherTimeRule,
)
from app.models.period import Period, PeriodTable, PeriodType
from app.models.reference_import import ReferenceImportBatch, ReferenceSchedulingRule
from app.models.semester import Semester
from app.models.timetable import ScheduleEntry, Timetable, TimetableStatus
from app.services.assignments import get_or_create_single_unit

ADAPTER_VERSION = "reference-v2"
TARGET_ACADEMIC_YEAR = 2026
TARGET_TERM = 1
TARGET_START = date(2026, 9, 1)
TARGET_END = date(2027, 1, 25)

CLASS_NAMES = {
    grade: [f"{grade}.{index}" for index in range(1, count + 1)]
    for grade, count in ((7, 5), (8, 4), (9, 3))
}

SUBJECT_ALIASES = {
    "道法": "道德与法治",
    "生物": "生物学",
    "信息": "信息科技",
    "信息技术": "信息科技",
    "体育": "体育与健康",
    "外语": "英语",
    "综合实践": "综合实践活动",
}

# Word 课时矩阵。复合的 4+1/3+2/3+1 在 _component_specs 中拆开。
WORD_TOTALS = {7: 35, 8: 35, 9: 35}
COMPONENTS: dict[int, list[tuple[str, str, int]]] = {
    7: [
        ("语文", "core", 6), ("数学", "core", 4), ("劳动", "math-plus", 1),
        ("英语", "core", 3), ("校本课程", "english-plus", 1), ("地方课程", "english-plus", 1),
        ("道德与法治", "core", 3), ("历史", "core", 2), ("生物学", "core", 3),
        ("地理", "core", 2), ("体育与健康", "physical", 4), ("音乐", "core", 1),
        ("美术", "core", 1), ("信息科技", "core", 1), ("班团队", "fixed-class-team", 1),
        ("国防体育", "national-defense", 1),
    ],
    8: [
        ("语文", "core", 5), ("数学", "core", 4), ("劳动", "math-plus", 1),
        ("英语", "core", 3), ("信息科技", "english-plus", 1), ("班团队", "english-plus", 1),
        ("道德与法治", "core", 3), ("历史", "core", 2), ("生物学", "core", 2),
        ("地理", "core", 2), ("物理", "core", 3), ("综合实践活动", "physics-plus", 1),
        ("体育与健康", "physical", 4), ("音乐", "core", 1), ("美术", "core", 1),
        ("国防体育", "national-defense", 1),
    ],
    9: [
        ("语文", "core", 6), ("数学", "core", 5), ("英语", "core", 3),
        ("劳动", "foreign-plus", 1), ("综合实践活动", "foreign-plus", 1),
        ("道德与法治", "core", 3), ("历史", "core", 2), ("物理", "core", 3),
        ("班团队", "physics-plus", 1), ("化学", "core", 3), ("体育与健康", "physical", 5),
        ("音乐", "physics-plus-independent", 1), ("美术", "chemistry-plus-independent", 1),
    ],
}

MAJOR_SUBJECTS = {"语文", "数学", "英语"}
ROOM_BY_SUBJECT = {
    "音乐": RoomType.special.value,
    "美术": RoomType.special.value,
    "信息科技": RoomType.special.value,
    "物理": RoomType.special.value,
    "化学": RoomType.special.value,
    "生物学": RoomType.special.value,
    "体育与健康": RoomType.outdoor.value,
    "国防体育": RoomType.outdoor.value,
}

RESEARCH_DAYS = {
    1: ("习字/书法", "综合实践活动"),
    3: ("英语", "生物学", "物理", "历史", "信息科技"),
    4: ("小学科学", "数学", "道德与法治", "政治", "地理", "化学", "美术", "技术"),
    5: ("语文", "心理", "音乐", "体育与健康"),
}

CLASS_RE = re.compile(r"([789])\s*[.．]\s*(\d+)")
CHINESE_RE = re.compile(r"[\u4e00-\u9fff]{2,8}")


class InvalidReferenceFile(ValueError):
    """文件不是预期的参考文件格式。"""


@dataclass(slots=True)
class AssignmentRow:
    class_name: str
    grade: int
    subject: str
    component: str
    periods: int
    teacher: str | None
    source_key: str
    notes: str = ""
    required_room_type: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "class_name": self.class_name,
            "grade": self.grade,
            "subject": self.subject,
            "component": self.component,
            "periods": self.periods,
            "teacher": self.teacher,
            "source_key": self.source_key,
            "notes": self.notes,
            "required_room_type": self.required_room_type,
        }


@dataclass(slots=True)
class RuleRow:
    scope: str
    grade: int | None
    subject_name: str | None
    weekday: int | None
    period_no: int | None
    rule_type: str
    source_key: str
    source_text: str
    enforced: bool

    def as_dict(self) -> dict[str, Any]:
        return {
            "scope": self.scope,
            "grade": self.grade,
            "subject_name": self.subject_name,
            "weekday": self.weekday,
            "period_no": self.period_no,
            "rule_type": self.rule_type,
            "source_key": self.source_key,
            "source_text": self.source_text,
            "enforced": self.enforced,
        }


@dataclass(slots=True)
class FixedRow:
    class_name: str
    subject: str
    teacher: str | None
    weekday: int
    period_no: int
    source_key: str
    notes: str = ""

    def as_dict(self) -> dict[str, Any]:
        return {
            "class_name": self.class_name,
            "subject": self.subject,
            "teacher": self.teacher,
            "weekday": self.weekday,
            "period_no": self.period_no,
            "source_key": self.source_key,
            "notes": self.notes,
        }


@dataclass(slots=True)
class ReferencePlan:
    semester_id: int
    fingerprint: str
    word_filename: str
    xlsx_filename: str
    word_sha256: str
    xlsx_sha256: str
    classes: list[dict[str, Any]] = field(default_factory=list)
    subjects: list[dict[str, Any]] = field(default_factory=list)
    teachers: dict[str, dict[str, Any]] = field(default_factory=dict)
    rooms: list[dict[str, Any]] = field(default_factory=list)
    assignments: list[AssignmentRow] = field(default_factory=list)
    rules: list[RuleRow] = field(default_factory=list)
    fixed_entries: list[FixedRow] = field(default_factory=list)
    errors: list[dict[str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    changes: list[dict[str, Any]] = field(default_factory=list)
    decisions: dict[str, Any] = field(default_factory=dict)

    @property
    def can_commit(self) -> bool:
        return not self.errors

    def as_dict(self) -> dict[str, Any]:
        counts = {
            "classes": len(self.classes),
            "teachers": len(self.teachers),
            "subjects": len(self.subjects),
            "rooms": len(self.rooms),
            "assignments": len(self.assignments),
            "rules": len(self.rules),
            "fixed_entries": len(self.fixed_entries),
        }
        return {
            "fingerprint": self.fingerprint,
            "adapter_version": ADAPTER_VERSION,
            "semester": {
                "id": self.semester_id,
                "academic_year": TARGET_ACADEMIC_YEAR,
                "term": TARGET_TERM,
                "start_date": TARGET_START.isoformat(),
                "end_date": TARGET_END.isoformat(),
            },
            "can_commit": self.can_commit,
            "has_changes": bool(self.changes),
            "counts": counts,
            "changes": self.changes,
            "assignments": [row.as_dict() for row in self.assignments],
            "rules": [row.as_dict() for row in self.rules],
            "fixed_entries": [row.as_dict() for row in self.fixed_entries],
            "errors": self.errors,
            "warnings": self.warnings,
            "source": {
                "word_filename": self.word_filename,
                "xlsx_filename": self.xlsx_filename,
                "word_sha256": self.word_sha256,
                "xlsx_sha256": self.xlsx_sha256,
            },
        }


@dataclass(slots=True)
class ParsedSource:
    mappings: dict[tuple[str, int], dict[str, list[str]]] = field(default_factory=dict)
    teachers: dict[str, dict[str, Any]] = field(default_factory=dict)
    teacher_subjects: dict[str, set[str]] = field(default_factory=dict)
    class_teachers: dict[str, str] = field(default_factory=dict)
    row_periods: dict[tuple[str, int], str] = field(default_factory=dict)
    raw_rows: list[str] = field(default_factory=list)
    source_rows: list[dict[str, Any]] = field(default_factory=list)


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _canonical_subject(name: str) -> str:
    return SUBJECT_ALIASES.get(name.strip(), name.strip())


def _read_docx_paragraphs(data: bytes) -> list[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("word/document.xml")
    except (KeyError, zipfile.BadZipFile) as exc:
        raise InvalidReferenceFile("Word 文件无法读取，请上传排课规则.docx") from exc
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    try:
        root = ET.fromstring(xml)
    except ET.ParseError as exc:
        raise InvalidReferenceFile("Word 文件内容损坏") from exc
    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:body/w:p", ns):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", ns)).strip()
        if text:
            paragraphs.append(text)
    if not paragraphs:
        raise InvalidReferenceFile("Word 文件没有可读取的排课规则")
    return paragraphs


def _parse_docx(data: bytes) -> list[str]:
    paragraphs = _read_docx_paragraphs(data)
    joined = "\n".join(paragraphs)
    required_markers = ("排课", "教研日")
    if not all(marker in joined for marker in required_markers):
        raise InvalidReferenceFile("Word 文件缺少排课规则或教研日内容")
    return paragraphs


def _grades(text: str) -> list[int]:
    result = []
    for label, grade in (("七", 7), ("八", 8), ("九", 9)):
        if label in text:
            result.append(grade)
    return result


def _expand_classes(matches: list[re.Match[str]], scopes: list[int]) -> list[str]:
    if not matches:
        return [name for grade in scopes for name in CLASS_NAMES[grade]]
    output: list[str] = []
    for match in matches:
        output.append(f"{match.group(1)}.{int(match.group(2))}")
    # 同年级区间（如 8.1-8.2）。
    for left, right in zip(matches, matches[1:], strict=False):
        if left.end() < right.start() and "-" not in left.string[left.end():right.start()]:
            continue
        if left.group(1) == right.group(1):
            grade = int(left.group(1))
            for index in range(int(left.group(2)), int(right.group(2)) + 1):
                output.append(f"{grade}.{index}")
        elif int(left.group(1)) < int(right.group(1)):
            for grade in range(int(left.group(1)), int(right.group(1)) + 1):
                output.extend(CLASS_NAMES.get(grade, []))
    return list(dict.fromkeys(output))


def _teacher_name(prefix: str) -> str | None:
    prefix = prefix.strip().lstrip("+ ")
    if not prefix:
        return None
    # 一个单元格可写“贾宁... + 徐盛 7.3”，取最后一个加号后的教师。
    if "+" in prefix:
        prefix = prefix.rsplit("+", 1)[-1].strip()
    matches = CHINESE_RE.findall(prefix)
    if not matches:
        return None
    ignored = {"班主任", "年级主管", "国防军事训练", "缺口", "一学期产假"}
    for candidate in matches:
        if (
            candidate not in ignored
            and "年级" not in candidate
            and not candidate.endswith("主管")
            and not candidate.endswith("班主任")
            and not candidate.startswith(("单周", "双周", "三周", "四周"))
        ):
            return candidate
    return None


def _segments(cell: str) -> list[str]:
    # 换行/分号才表示不同教师；括号里的单双周说明不能拆开。
    parts: list[str] = []
    current: list[str] = []
    depth = 0
    opening = {"(", "（"}
    closing = {")": "(", "）": "（",
    }
    for char in cell:
        if char in opening:
            depth += 1
        elif char in closing and depth:
            depth -= 1
        if depth == 0 and char in {"\n", ";", "；"}:
            part = "".join(current).strip()
            if part:
                parts.append(part)
            current = []
            continue
        current.append(char)
    part = "".join(current).strip()
    if part:
        parts.append(part)
    return parts


def _teacher_chunks(segment: str) -> list[tuple[str | None, list[re.Match[str]]]]:
    """拆分一个单元格中用加号连接的多位教师/备注。

    右侧没有新的教师名时（例如 ``贺温雁 9.1 + 7.5（跨头）``），
    其班号仍归给左侧教师；右侧出现教师名时（例如 ``贾宁 7.1 + 徐盛 7.3``）
    才切换教师。
    """
    chunks = [part.strip() for part in re.split(r"\s*\+\s*", segment) if part.strip()]
    result: list[tuple[str | None, list[re.Match[str]]]] = []
    current: str | None = None
    for chunk in chunks:
        matches = list(CLASS_RE.finditer(chunk))
        candidate = _teacher_name(chunk[: matches[0].start()] if matches else chunk)
        if candidate is not None:
            current = candidate
        result.append((current, matches))
    return result


def _parse_xlsx(data: bytes, *, include_word_history_fallback: bool = True) -> ParsedSource:
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001
        raise InvalidReferenceFile("Excel 文件无法读取，请上传教师安排8.24.xlsx") from exc
    sheet = workbook.active
    if sheet is None:
        raise InvalidReferenceFile("Excel 文件没有工作表")
    parsed = ParsedSource()
    current_subject: str | None = None
    for row_no in range(4, min(sheet.max_row, 42) + 1):
        subject_value = sheet.cell(row_no, 1).value
        if subject_value:
            current_subject = _canonical_subject(str(subject_value))
        grade_text = str(sheet.cell(row_no, 3).value or "")
        row_values = [sheet.cell(row_no, column).value for column in range(1, 12)]
        if any(value is not None for value in row_values):
            parsed.source_rows.append(
                {
                    "row": row_no,
                    "subject": current_subject,
                    "grade": grade_text,
                    "weekly_periods": sheet.cell(row_no, 4).value,
                    "fee": sheet.cell(row_no, 2).value,
                    "existing_staff": sheet.cell(row_no, 5).value,
                    "shortage_staff": sheet.cell(row_no, 6).value,
                    "assignment_grade": sheet.cell(row_no, 7).value,
                    "assignment_cells": [
                        sheet.cell(row_no, column).value for column in range(8, 12)
                    ],
                    "need": sheet.cell(row_no, 11).value,
                }
            )
        if not current_subject or current_subject == "校医":
            # 校医信息只保留为被排除的来源记录。
            continue
        scopes = _grades(grade_text)
        if not scopes:
            scopes = [7, 8, 9] if current_subject in {"美术", "音乐", "心理"} else []
        parsed.row_periods[(current_subject, scopes[0] if len(scopes) == 1 else 0)] = str(
            sheet.cell(row_no, 4).value or ""
        )
        for column in range(8, 12):
            value = sheet.cell(row_no, column).value
            if value is None:
                continue
            cell_text = str(value).strip()
            if not cell_text:
                continue
            parsed.raw_rows.append(f"{row_no}:{current_subject}:{cell_text}")
            for segment in _segments(cell_text):
                for name, matches in _teacher_chunks(segment):
                    if name is None:
                        continue
                    notes = segment
                    if "金铭" in name:
                        continue
                    parsed.teacher_subjects.setdefault(name, set()).add(current_subject)
                    info = parsed.teachers.setdefault(
                        name,
                        {"notes": [], "external": False, "admin": False, "maternity": False},
                    )
                    if notes not in info["notes"]:
                        info["notes"].append(notes)
                    info["external"] |= "外聘" in notes
                    info["admin"] |= "行政" in notes
                    info["maternity"] |= "产假" in notes
                    if not matches:
                        continue
                    classes = _expand_classes(matches, scopes)
                    for class_name in classes:
                        grade = int(class_name.split(".", 1)[0])
                        if scopes and grade not in scopes and current_subject not in {"历史", "地理"}:
                            continue
                        mapping = parsed.mappings.setdefault((current_subject, grade), {}).setdefault(
                            class_name, []
                        )
                        if name not in mapping:
                            mapping.append(name)
                        if "班主任" in notes and class_name not in parsed.class_teachers:
                            parsed.class_teachers[class_name] = name
    if include_word_history_fallback:
        # Word 明确要求七年级历史，Excel 行仅标注八、九年级；沿用同一行的杨紫文映射。
        history = parsed.mappings.setdefault(("历史", 7), {})
        for class_name in CLASS_NAMES[7]:
            history.setdefault(class_name, ["杨紫文"])
    return parsed


def _db_snapshot(db: Session, semester_id: int) -> dict[str, Any]:
    assignments = db.scalars(select(CourseAssignment).where(CourseAssignment.semester_id == semester_id)).all()
    teachers = db.scalars(select(Teacher).where(Teacher.semester_id == semester_id)).all()
    classes = db.scalars(select(ClassUnit).where(ClassUnit.semester_id == semester_id)).all()
    return {
        "assignments": sorted(a.reference_source_key or f"manual:{a.id}" for a in assignments),
        "teachers": sorted(t.name for t in teachers),
        "classes": sorted(c.name for c in classes),
    }


def _source_key(class_name: str, subject: str, component: str, teacher: str | None) -> str:
    value = teacher or "unassigned"
    return f"{ADAPTER_VERSION}:{class_name}:{subject}:{component}:{value}"[:160]


def _mapping_teacher(source: ParsedSource, subject: str, grade: int, class_name: str) -> str | None:
    names = source.mappings.get((subject, grade), {}).get(class_name, [])
    if names:
        return names[0]
    # Word 明确覆盖七、八、九年级的音乐/美术；Excel 的区间只写到 8.4，
    # 因此九年级沿用同一科目的首位教师，而不是生成缺口。
    if subject in {"音乐", "美术"}:
        for (mapped_subject, _), grade_map in source.mappings.items():
            if mapped_subject == subject:
                for values in grade_map.values():
                    if values:
                        return values[0]
    return None


def _component_specs(source: ParsedSource, warnings: list[str]) -> list[AssignmentRow]:
    rows: list[AssignmentRow] = []
    independent_subjects = {"音乐", "美术", "信息科技", "历史"}
    for grade, components in COMPONENTS.items():
        for class_name in CLASS_NAMES[grade]:
            for subject, component, periods in components:
                if component == "national-defense":
                    teacher = "刘锴"
                elif component == "fixed-class-team":
                    teacher = source.class_teachers.get(class_name)
                elif component == "physics-plus-independent":
                    # 音乐独立行优先；物理/班团队本身在其它组件中单独生成。
                    teacher = _mapping_teacher(source, subject, grade, class_name)
                elif component == "chemistry-plus-independent":
                    teacher = _mapping_teacher(source, subject, grade, class_name)
                else:
                    source_subject = subject
                    if subject in {"劳动", "校本课程", "地方课程", "班团队", "综合实践活动"} or (
                        subject == "信息科技" and grade == 8
                    ):
                        source_subject = {
                            "劳动": "数学" if grade in {7, 8} else "英语",
                            "校本课程": "英语",
                            "地方课程": "英语",
                            "班团队": "英语" if grade == 8 else ("物理" if grade == 9 else "班团队"),
                            "综合实践活动": "物理" if grade == 8 else "英语",
                            "信息科技": "英语",
                        }[subject]
                    teacher = _mapping_teacher(source, source_subject, grade, class_name)
                    if subject == "班团队" and grade == 7:
                        teacher = source.class_teachers.get(class_name)
                if subject in independent_subjects:
                    teacher = _mapping_teacher(source, subject, grade, class_name) or teacher
                if teacher is None and subject not in {"班团队"}:
                    warnings.append(f"{class_name} 的{subject}缺少教师映射，需要在预览中人工指定")
                rows.append(
                    AssignmentRow(
                        class_name=class_name,
                        grade=grade,
                        subject=subject,
                        component=component,
                        periods=periods,
                        teacher=teacher,
                        source_key=_source_key(class_name, subject, component, teacher),
                        notes="Word 课时权威；Excel 教师映射" if teacher else "缺少教师映射",
                        required_room_type=ROOM_BY_SUBJECT.get(subject),
                    )
                )
    # 独立的信息科技行覆盖七年级英语复合规则中的同名组件，不增加第二份课时。
    for row in rows:
        if row.subject == "信息科技" and row.grade == 7:
            row.teacher = _mapping_teacher(source, "信息科技", row.grade, row.class_name)
            row.source_key = _source_key(row.class_name, row.subject, row.component, row.teacher)
    # 心理来源保留教师和警告，但不创建任务，不增加班级总课时。
    if "孙佳伟" in source.teachers:
        warnings.append("心理的双周/三周/四周周期已忽略，按数学兼课处理，不生成额外心理课时")
    # 国防体育在七、八年级占用原体育总量中的一节；任务拆成 4+1，
    # 因此不再减少普通体育的课时字段。
    for grade in (7, 8, 9):
        for class_name in CLASS_NAMES[grade]:
            physical = [row for row in rows if row.class_name == class_name and row.subject == "体育与健康"]
            if not physical:
                continue
            target = physical[-1]
            if grade in {7, 8}:
                target.notes += "；国防体育替代1节"
            # 七年级 7.5 的 Excel 行明确拆成刘锴 2 节、周嘉辉 3 节；
            # 国防体育替代其中一节后保留 1+3 的普通体育分工。
            if grade == 7 and class_name == "7.5":
                names = source.mappings.get(("体育与健康", 7), {}).get(class_name, [])
                if len(names) > 1:
                    target.teacher = names[0]
                    target.periods = 1
                    target.source_key = _source_key(target.class_name, target.subject, target.component, target.teacher)
                    rows.append(
                        AssignmentRow(
                            class_name=class_name,
                            grade=grade,
                            subject="体育与健康",
                            component="physical-split",
                            periods=3,
                            teacher=names[1],
                            source_key=_source_key(class_name, "体育与健康", "physical-split", names[1]),
                            notes="Excel 明确的 7.5 班体育分工",
                            required_room_type=ROOM_BY_SUBJECT["体育与健康"],
                        )
                    )
    return rows


def _rule_specs(source: ParsedSource) -> list[RuleRow]:
    rules: list[RuleRow] = []
    for weekday, subjects in RESEARCH_DAYS.items():
        for subject in subjects:
            rules.append(RuleRow("subject", None, _canonical_subject(subject), weekday, None, "avoid", f"research:{weekday}:{subject}", f"周{weekday}为{subject}研究日，尤其避免上午", False))
    rules.extend(
        [
            RuleRow("global", None, None, 1, 1, "unavailable", "global:mon-p1", "全校周一上午第1节不排行政会议", True),
            RuleRow("subject", None, "体育与健康", None, 3, "window", "subject:pe-from-p3", "体育一般从第3节开始", False),
            RuleRow("subject", None, "语文", None, None, "prefer", "subject:major-morning", "语文、数学、英语优先上午前3节", False),
            RuleRow("room", None, None, None, None, "capacity", "room:special-one-each", "音乐、美术、机房、理化生实验室各一间，按资源冲突校验", True),
            RuleRow("national-defense", 7, "国防体育", None, 3, "window", "national-defense:grade7:tue-fri-p3-p4", "七年级国防体育安排在周二至周五第3/4节候选时段", False),
            RuleRow("national-defense", 8, "国防体育", None, 3, "window", "national-defense:grade8:tue-fri-p3-p4", "八年级国防体育安排在周二至周五第3/4节候选时段", False),
        ]
    )
    return rules


def _teacher_rule_rows(teachers: dict[str, dict[str, Any]]) -> list[RuleRow]:
    rows: list[RuleRow] = []
    for name in teachers:
        rows.append(RuleRow("teacher", None, None, 1, 1, "unavailable", f"teacher:{name}:mon-p1", f"{name}遵守周一上午第1节全校会议", True))
    for name in ("贾宁", "孙佳伟", "裴树丹", "宋广莲", "王少凤"):
        if name in teachers:
            rows.append(RuleRow("teacher", None, None, 1, 6, "unavailable", f"teacher:{name}:class-meeting", "班主任会议固定占用周一下午第2节", True))
    return rows


def _apply_teacher_time_rules(db: Session, teacher: Teacher) -> None:
    existing = {(r.weekday, r.period_no): r for r in teacher.time_rules}
    wanted: dict[tuple[int, int], str] = {}

    def add(weekday: int, period: int, rule: str) -> None:
        wanted.setdefault((weekday, period), rule)

    for weekday in range(1, 6):
        add(weekday, 1, "unavailable")
    if teacher.name in {"贾宁", "孙佳伟", "裴树丹", "宋广莲", "王少凤"}:
        add(1, 6, "unavailable")
    if teacher.name == "朱振华":
        for weekday in range(1, 6):
            for period in range(1, 9):
                if period not in {2, 3}:
                    add(weekday, period, "unavailable")
    if teacher.name == "唐延艳":
        for weekday in range(1, 6):
            add(weekday, 1, "unavailable")
    if teacher.name == "崔洪刚":
        for period in range(5, 9):
            add(1, period, "unavailable")
        for period in range(1, 9):
            add(4, period, "unavailable")
        for weekday in range(1, 6):
            add(weekday, 8, "unavailable")
    if teacher.name == "朱峻":
        for period in range(5, 9):
            add(1, period, "unavailable")
    if teacher.name == "张灿":
        for period in range(5, 9):
            add(3, period, "unavailable")
    if teacher.name == "孟召磊":
        for weekday in range(1, 6):
            add(weekday, 4, "avoid")
    if teacher.name == "路兆宇":
        for weekday in range(1, 6):
            add(weekday, 4, "avoid")
        for period in range(1, 5):
            add(5, period, "avoid")
    for (weekday, period), rule_type in wanted.items():
        if (weekday, period) in existing:
            continue
        db.add(TeacherTimeRule(teacher_id=teacher.id, weekday=weekday, period_no=period, rule_type=rule_type))


def _ensure_period_table(db: Session, semester_id: int) -> PeriodTable:
    tables = list(db.scalars(select(PeriodTable).where(PeriodTable.semester_id == semester_id).order_by(PeriodTable.id)))
    table = next((item for item in tables if item.is_default), tables[0] if tables else None)
    if table is None:
        table = PeriodTable(semester_id=semester_id, name="初中五天八节", num_weekdays=5, is_default=True)
        db.add(table)
        db.flush()
    existing = {(p.weekday, p.period_no) for p in table.periods}
    for weekday in range(1, 6):
        for period_no in range(1, 9):
            if (weekday, period_no) not in existing:
                db.add(Period(period_table_id=table.id, weekday=weekday, period_no=period_no, name=f"第{period_no}节", type=PeriodType.regular.value))
    db.flush()
    return table


def build_plan(
    db: Session,
    semester_id: int,
    word_bytes: bytes,
    xlsx_bytes: bytes,
    *,
    word_filename: str = "排课规则.docx",
    xlsx_filename: str = "教师安排8.24.xlsx",
    decisions: dict[str, Any] | None = None,
) -> ReferencePlan:
    paragraphs = _parse_docx(word_bytes)
    source = _parse_xlsx(xlsx_bytes)
    plan = ReferencePlan(
        semester_id=semester_id,
        fingerprint="",
        word_filename=word_filename,
        xlsx_filename=xlsx_filename,
        word_sha256=_sha(word_bytes),
        xlsx_sha256=_sha(xlsx_bytes),
        decisions=decisions or {},
    )
    semester = db.get(Semester, semester_id)
    if semester is None:
        plan.errors.append({"code": "semester_not_found", "message": "找不到目标学期", "source": str(semester_id)})
    else:
        if (semester.academic_year, semester.term) != (TARGET_ACADEMIC_YEAR, TARGET_TERM):
            plan.errors.append({"code": "semester_mismatch", "message": "参考文件适配器只接受 2026-2027 第1学期", "source": semester.label})
        if semester.start_date != TARGET_START or semester.end_date != TARGET_END:
            plan.errors.append({"code": "semester_dates_mismatch", "message": "目标学期起止日期必须为 2026-09-01 至 2027-01-25", "source": f"{semester.start_date} - {semester.end_date}"})
    plan.classes = [
        {
            "name": name,
            "grade": grade,
            "track": ClassTrack.junior_high.value,
            "homeroom_teacher": source.class_teachers.get(name),
        }
        for grade in (7, 8, 9)
        for name in CLASS_NAMES[grade]
    ]
    subject_names = sorted({subject for values in COMPONENTS.values() for subject, _, _ in values} | {"心理健康教育"})
    plan.subjects = [{"name": name, "is_major": name in MAJOR_SUBJECTS, "required_room_type": ROOM_BY_SUBJECT.get(name)} for name in subject_names]
    plan.teachers = {name: {**info, "notes": "；".join(info["notes"])} for name, info in source.teachers.items() if name != "金铭"}
    # 规则中的教师即使 Excel 没有分配课，也需要进入可排教师集合。
    for name in ("朱振华", "唐延艳", "崔洪刚", "朱峻", "张灿", "孟召磊", "路兆宇", "贾宁", "孙佳伟", "裴树丹", "宋广莲", "王少凤", "刘锴"):
        plan.teachers.setdefault(name, {"notes": "Word 规则教师", "external": False, "admin": False, "maternity": False})
    plan.rooms = [
        {"name": "音乐教室", "room_type": RoomType.special.value},
        {"name": "美术教室", "room_type": RoomType.special.value},
        {"name": "机房", "room_type": RoomType.special.value},
        {"name": "理化生实验室", "room_type": RoomType.special.value},
        {"name": "操场", "room_type": RoomType.outdoor.value},
    ]
    plan.assignments = _component_specs(source, plan.warnings)
    plan.rules = _rule_specs(source) + _teacher_rule_rows(plan.teachers)
    plan.fixed_entries = [
        FixedRow(class_name, "班团队", source.class_teachers.get(class_name), 2, 3, f"fixed:{class_name}:class-team:tue-p3", "Word 固定：七年级班团队周二第3节")
        for class_name in CLASS_NAMES[7]
    ]
    plan.warnings.append("Excel 中‘课时费’按费用处理，未作为每周课时；‘缺口’仅作为提示，不自动生成教师")
    plan.warnings.append("产假教师按正常在岗教师参与排课；校医金铭被排除，不创建教学任务")
    plan.warnings.append("Word 对体育、信息和七年级历史的范围/冲突优先于 Excel 表头")
    plan.warnings.append("九年级按本次导入决策覆盖 Word 原文 34 节，按 35 节生成")
    for class_name in CLASS_NAMES[7]:
        if class_name not in source.class_teachers:
            plan.warnings.append(f"{class_name} 未能从来源文本确定班主任，班团队任务将等待人工指定教师")
    if not source.mappings.get(("信息科技", 8)):
        plan.warnings.append("Excel 未提供八年级信息科技独立教师；按英语 3+2 的复合任务由英语教师承担")
    totals: dict[str, int] = {name: 0 for name in (name for grade in CLASS_NAMES.values() for name in grade)}
    for row in plan.assignments:
        totals[row.class_name] += row.periods
    for class_name, total in totals.items():
        if total != WORD_TOTALS[int(class_name.split(".", 1)[0])]:
            plan.errors.append({"code": "class_total_mismatch", "message": f"{class_name} 归一化后为 {total} 节，目标为 35 节", "source": class_name})
    existing_keys = {key for key in _db_snapshot(db, semester_id)["assignments"] if key.startswith(ADAPTER_VERSION)}
    incoming_keys = {row.source_key for row in plan.assignments}
    for row in plan.assignments:
        plan.changes.append({"entity": "course_assignment", "status": "unchanged" if row.source_key in existing_keys else "new", "identity": row.source_key, "details": row.as_dict()})
    for class_item in plan.classes:
        plan.changes.append({"entity": "class", "status": "existing" if class_item["name"] in _db_snapshot(db, semester_id)["classes"] else "new", "identity": class_item["name"]})
    disappeared = sorted(existing_keys - incoming_keys)
    if disappeared:
        plan.warnings.append(f"检测到 {len(disappeared)} 个此前导入但本次文件消失的教学任务；仅提示，不删除")
    fingerprint_payload = {
        "adapter": ADAPTER_VERSION,
        "semester": semester_id,
        "word": plan.word_sha256,
        "xlsx": plan.xlsx_sha256,
        "decisions": plan.decisions,
        "database": _db_snapshot(db, semester_id),
        "paragraphs": paragraphs,
    }
    plan.fingerprint = hashlib.sha256(json.dumps(fingerprint_payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return plan


def _find_or_create_subject(db: Session, semester_id: int, item: dict[str, Any]) -> Subject:
    subject = db.scalar(select(Subject).where(Subject.semester_id == semester_id, Subject.name == item["name"]))
    if subject is None:
        subject = Subject(semester_id=semester_id, name=item["name"], is_major=item["is_major"], required_room_type=item.get("required_room_type"))
        db.add(subject)
        db.flush()
    return subject


def apply_plan(db: Session, plan: ReferencePlan) -> dict[str, Any]:
    if plan.errors:
        raise ValueError("参考文件导入计划仍有冲突")
    duplicate = db.scalar(
        select(ReferenceImportBatch).where(
            ReferenceImportBatch.semester_id == plan.semester_id,
            ReferenceImportBatch.word_sha256 == plan.word_sha256,
            ReferenceImportBatch.xlsx_sha256 == plan.xlsx_sha256,
        )
    )
    if duplicate is not None and (duplicate.decisions or {}) == (plan.decisions or {}):
        # File hashes alone are not enough: an operator may have removed an
        # imported task after the previous batch.  Only skip the write when
        # every source identity in the current plan is still present.
        incoming_keys = {row.source_key for row in plan.assignments}
        existing_keys = set(
            db.scalars(
                select(CourseAssignment.reference_source_key).where(
                    CourseAssignment.semester_id == plan.semester_id,
                    CourseAssignment.reference_source_key.in_(incoming_keys),
                )
            ).all()
        )
        if existing_keys == incoming_keys:
            return {"batch_id": duplicate.id, "timetable_id": duplicate.summary.get("timetable_id"), "created": {}, "unchanged": {"batch": 1}, "warnings": plan.warnings, "idempotent": True}
    semester = db.get(Semester, plan.semester_id)
    if semester is None:
        raise ValueError("目标学期不存在")
    subjects = {item["name"]: _find_or_create_subject(db, plan.semester_id, item) for item in plan.subjects}
    teachers: dict[str, Teacher] = {}
    for name, info in plan.teachers.items():
        teacher = db.scalar(select(Teacher).where(Teacher.semester_id == plan.semester_id, Teacher.name == name))
        if teacher is None:
            teacher = Teacher(semester_id=plan.semester_id, name=name, base_periods=0, is_external=bool(info.get("external")), is_active=True, admin_title="行政" if info.get("admin") else None)
            db.add(teacher)
            db.flush()
        elif info.get("maternity") and not teacher.is_active:
            plan.warnings.append(f"{name} 已存在但被标记为停用，请人工确认后再排课")
        teachers[name] = teacher
        for subject in subjects.values():
            if subject.name in {row.subject for row in plan.assignments if row.teacher == name} and subject not in teacher.subjects:
                teacher.subjects.append(subject)
        _apply_teacher_time_rules(db, teacher)
    classes: dict[str, ClassUnit] = {}
    table = _ensure_period_table(db, plan.semester_id)
    for item in plan.classes:
        class_unit = db.scalar(select(ClassUnit).where(ClassUnit.semester_id == plan.semester_id, ClassUnit.name == item["name"]))
        if class_unit is None:
            homeroom_teacher = teachers.get(item.get("homeroom_teacher"))
            class_unit = ClassUnit(
                semester_id=plan.semester_id,
                grade=item["grade"],
                name=item["name"],
                track=item["track"],
                homeroom_teacher_id=homeroom_teacher.id if homeroom_teacher else None,
                period_table_id=table.id,
            )
            db.add(class_unit)
            db.flush()
        elif class_unit.period_table_id is None:
            class_unit.period_table_id = table.id
        # Keep an existing manual homeroom assignment; fill only an empty field.
        if class_unit.homeroom_teacher_id is None:
            homeroom_teacher = teachers.get(item.get("homeroom_teacher"))
            if homeroom_teacher is not None:
                class_unit.homeroom_teacher_id = homeroom_teacher.id
        classes[item["name"]] = class_unit
    rooms: dict[str, Room] = {}
    for item in plan.rooms:
        room = db.scalar(select(Room).where(Room.semester_id == plan.semester_id, Room.name == item["name"]))
        if room is None:
            room = Room(semester_id=plan.semester_id, name=item["name"], room_type=item["room_type"])
            db.add(room)
            db.flush()
        rooms[item["name"]] = room
    db.flush()
    assignments: dict[str, CourseAssignment] = {}
    created = 0
    unchanged = 0
    for row in plan.assignments:
        assignment = db.scalar(select(CourseAssignment).where(CourseAssignment.semester_id == plan.semester_id, CourseAssignment.reference_source_key == row.source_key))
        if assignment is not None:
            assignments[row.source_key] = assignment
            unchanged += 1
            continue
        unit = get_or_create_single_unit(db, classes[row.class_name])
        assignment = CourseAssignment(semester_id=plan.semester_id, scheduling_unit_id=unit.id, subject_id=subjects[row.subject].id, periods_per_week=row.periods, required_room_type=row.required_room_type, reference_source_key=row.source_key)
        db.add(assignment)
        db.flush()
        if row.teacher and row.teacher in teachers:
            db.add(AssignmentTeacher(course_assignment_id=assignment.id, teacher_id=teachers[row.teacher].id, is_lead=True))
        assignments[row.source_key] = assignment
        created += 1
    for rule in plan.rules:
        existing = db.scalar(select(ReferenceSchedulingRule).where(ReferenceSchedulingRule.semester_id == plan.semester_id, ReferenceSchedulingRule.source_key == rule.source_key))
        if existing is None:
            db.add(ReferenceSchedulingRule(semester_id=plan.semester_id, scope=rule.scope, grade=rule.grade, subject_name=rule.subject_name, weekday=rule.weekday, period_no=rule.period_no, rule_type=rule.rule_type, source_key=rule.source_key, source_text=rule.source_text, enforced=rule.enforced))
    timetable_name = f"{semester.label}·参考文件草稿"
    timetable = db.scalar(select(Timetable).where(Timetable.semester_id == plan.semester_id, Timetable.name == timetable_name, Timetable.status == TimetableStatus.draft.value))
    if timetable is None:
        from app.services.scheduling_rules import active_revision_id

        timetable = Timetable(
            semester_id=plan.semester_id,
            rule_revision_id=active_revision_id(db, plan.semester_id),
            name=timetable_name,
            status=TimetableStatus.draft.value,
        )
        db.add(timetable)
        db.flush()
    for fixed in plan.fixed_entries:
        source_key = _source_key(fixed.class_name, fixed.subject, "fixed-class-team", fixed.teacher)
        assignment = assignments.get(source_key)
        if assignment is None:
            assignment = db.scalar(select(CourseAssignment).where(CourseAssignment.semester_id == plan.semester_id, CourseAssignment.reference_source_key == source_key))
        if assignment is None:
            unit = get_or_create_single_unit(db, classes[fixed.class_name])
            assignment = CourseAssignment(semester_id=plan.semester_id, scheduling_unit_id=unit.id, subject_id=subjects[fixed.subject].id, periods_per_week=1, reference_source_key=source_key)
            db.add(assignment)
            db.flush()
            if fixed.teacher and fixed.teacher in teachers:
                db.add(AssignmentTeacher(course_assignment_id=assignment.id, teacher_id=teachers[fixed.teacher].id, is_lead=True))
        existing_entry = db.scalar(select(ScheduleEntry).where(ScheduleEntry.timetable_id == timetable.id, ScheduleEntry.course_assignment_id == assignment.id, ScheduleEntry.weekday == fixed.weekday, ScheduleEntry.period_no == fixed.period_no))
        if existing_entry is None:
            db.add(ScheduleEntry(timetable_id=timetable.id, course_assignment_id=assignment.id, weekday=fixed.weekday, period_no=fixed.period_no, span=1, locked=True))
    batch = ReferenceImportBatch(semester_id=plan.semester_id, fingerprint=plan.fingerprint, adapter_version=ADAPTER_VERSION, word_filename=plan.word_filename, xlsx_filename=plan.xlsx_filename, word_sha256=plan.word_sha256, xlsx_sha256=plan.xlsx_sha256, decisions=plan.decisions, summary={"timetable_id": timetable.id, "created_assignments": created, "unchanged_assignments": unchanged})
    db.add(batch)
    db.flush()
    return {"batch_id": batch.id, "timetable_id": timetable.id, "created": {"assignments": created}, "unchanged": {"assignments": unchanged}, "warnings": plan.warnings, "idempotent": False}
